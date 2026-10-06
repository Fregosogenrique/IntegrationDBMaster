"""Homologación del Cíclico integrado bajo el estándar G1522: un Cíclico por marca.

Toma la hoja Productos de Ciclico_1522_investigado.xlsx (catálogo NetSuite integrado con
Shopify Stetson, Shopify Western Brothers y Odoo), le aplica las reglas de
homologar_ariat_g1522.py y de la hoja Criterios, y escribe un archivo por marca con la
misma estructura que Ciclico_1522_Stetson.xlsx (plantilla): las 12 hojas, las 44 columnas
de Productos, sus fórmulas, validaciones, formatos condicionales y tablas.

Secuencia (Estructura_datos_validaciones_G1522.md, «Secuencia de validación y carga»):
  1. Preparar   - leer el Cíclico; las correcciones hechas a mano en los
                  Ciclico_1522_<Marca>.xlsx existentes mandan sobre él. Ariat suma la base
                  homologada (Base_Unificada_Ariat_G1522.xlsx), que guarda la evidencia de
                  su tienda Shopify.
  2. Marca      - Marca principal; si es Multimarca, WB Licencia; si no, WB Marca; sin
                  marca: la del mismo estilo, el nombre en la descripción o el prefijo del
                  código.
  3. Normalizar - Criterios (homologar_ariat_g1522.normalizar) más los casos propios del
                  Cíclico (país, licencia, CORE, Pieza, ceros iniciales, coma decimal,
                  acentos en materiales, estilo desde el SKU Karman), la investigación en
                  línea (Investigacion_web_G1522.csv) y Criterios › 3, 7 y 8 (talla de
                  EE. UU., Fit, Silueta).
  4. Integrar   - los registros que son el mismo producto se funden en uno; los originales
                  quedan en la hoja Revisión Duplicados.
  5. Validar    - controles G1522 y los cuatro validadores (NetSuite, Odoo, Shopify WB,
                  Shopify Stetson).
  6. Resolver   - cada cambio queda en «Notas de enriquecimiento / revisión» de su fila y
                  cada fuente web en «Fuentes de consulta».

  7. Inventario  - existencias de Productos y Stock por Ubicación al corte NetSuite 06/10,
                  más las hojas «Inventario» y «KPIs Inventario» (inventario_g1522.py).

Uso:    python3 homologar_ciclico_g1522.py [Marca ...]
Salida: Ciclico_1522_<Marca>.xlsx en esta carpeta (Montana West y Wrangler juntos).
"""

import io
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd

import homologar_ariat_g1522 as g
import inventario_g1522 as inv

BASE = Path(__file__).resolve().parent
CICLICO = BASE / "Ciclico_1522_investigado.xlsx"
ARIAT = BASE / "Base_Unificada_Ariat_G1522.xlsx"
PLANTILLA = BASE / "Ciclico_1522_Stetson.xlsx"
WEB = BASE / "Investigacion_web_G1522.csv"
FECHA = "06/10/2026"
# Arriba de este número de variantes, la validación de Productos (AA/AB) y la fila de
# Stock por Ubicación se escriben ya calculadas: sus fórmulas comparan cada fila contra
# toda la columna y con decenas de miles de filas saturan el navegador.
LIMITE_FORMULAS = 10000
HOJAS_CICLICO = ["Escaneo Diario", "Historial de Escaneos", "Escaneos", "Configuración", "Productos",
                 "Productos_Resumen", "Criterios", "Revisión Duplicados", "Stock por Ubicación",
                 "Bodega 25", "Bodega 43", "Resumen Integración"]

ATRIBUTOS = g.ATRIBUTOS
CAMPOS_FUSION = [c for c in ATRIBUTOS if c not in ("Código de barras", "Identificador interno")]

# Marcas que se trabajan en un mismo archivo.
GRUPOS = {"Montana West": "Montana West + Wrangler", "Wrangler": "Montana West + Wrangler"}

g.TEXTOS_SUSTITUTOS |= {"(blank)", "(Blank)", "0.0"}

PAISES = {"mexico": "México", "mx": "México", "mex": "México",
          "usa": "Estados Unidos", "us": "Estados Unidos", "eua": "Estados Unidos",
          "eeuu": "Estados Unidos", "estados unidos": "Estados Unidos",
          "china": "China", "india": "India", "vietnam": "Vietnam"}
# Materiales que suelen llegar sin acento; se corrige solo el acento, sin cambiar mayúsculas.
MATERIAL_ESCRITURA = {"algodon": "algodón", "poliester": "poliéster", "rayon": "rayón",
                      "acrilico": "acrílico"}

# Prioridad del registro que se conserva al fundir duplicados (árbol de decisión,
# Propuesta_Limpieza_Catalogo_Shopify_G1522.md §7): estructura y existencias primero.
PRIORIDAD_ORIGEN = [("Catálogo existente", 50), ("NetSuite con existencia", 46),
                    ("NetSuite —", 45), ("NetSuite sin existencia", 40),
                    ("Hoja a Revisar — artículo", 30), ("Shopify Stetson", 22),
                    ("Shopify Western", 22), ("Odoo", 20), ("Shopify Ariat", 15),
                    ("Carga Hoja", 10)]

NUMERICAS = ["Stock Sistema NetSuite", "Disponible NetSuite", "Stock Shopify Stetson México",
             "Stock Shopify Western Brothers", "Stock Odoo Universal Unique Brands"]
TEXTOS_UNION = ["Ubicaciones con stock (NetSuite)", "Ubicaciones con existencia negativa",
                "📍 Ubicación en plataformas (Plataforma + Marca)",
                "Ubicaciones con stock (plataformas)", "Fuentes de consulta"]
SHOPIFY_ARIAT = ["Shopify Handle", "Shopify ID producto", "Shopify ID variante", "Shopify Estatus"]
# Marcas de lo que este script agrega a las columnas de texto del Cíclico; al releer un
# Cíclico por marca ya generado se quitan para no duplicarlas.
MARCA_NOTA = "Homologación G1522"
MARCA_WEB = "Investigación en línea"


# --------------------------------------------------------------------------
# Bitácora
# --------------------------------------------------------------------------

class Revision(g.Revision):
    """La bitácora de Ariat con la fila de origen, para repartirla por libro."""

    COLUMNAS = g.Revision.COLUMNAS[:3] + ["WB Marca"] + g.Revision.COLUMNAS[3:]

    def add(self, fila, campo, origen, final, tipo, accion, evidencia=""):
        self.filas.append([fila.get("_uid", -1), fila.get("Código de barras", ""),
                           fila.get("WB SKU", ""), fila.get("Identificador interno", ""),
                           fila.get("WB Marca", ""), campo, origen, final, tipo, accion,
                           evidencia, ""])

    def tabla(self):
        return pd.DataFrame(self.filas, columns=["_uid"] + self.COLUMNAS)


def cambiar(df, i, campo, nuevo, rev, tipo, accion, evidencia=""):
    g.cambiar(df, i, campo, nuevo, rev, tipo, accion, evidencia)


# --------------------------------------------------------------------------
# 1. Preparar
# --------------------------------------------------------------------------

def leer_hoja(xl, hoja):
    df = pd.read_excel(xl, sheet_name=hoja, dtype=str).fillna("")
    df.columns = [str(c).strip() for c in df.columns]
    return df.apply(lambda s: s.str.strip())


def leer_ciclico():
    xl = pd.ExcelFile(CICLICO, engine="calamine")
    productos = leer_hoja(xl, "Productos")
    productos = productos[(productos[ATRIBUTOS] != "").any(axis=1)].reset_index(drop=True)
    productos = aplicar_ciclicos_por_marca(productos)
    hojas = {h: leer_hoja(xl, h) for h in ("Stock por Ubicación", "Escaneos",
                                           "Revisión Duplicados")}
    criterios = pd.read_excel(xl, sheet_name="Criterios", dtype=str, header=None).fillna("")
    criterios = criterios.apply(lambda s: s.str.strip())
    return productos, hojas, leer_criterios(criterios)


def llave_fila(df):
    """Llave para reconocer la misma fila entre el Cíclico y un Cíclico por marca."""
    llave = df["Código de barras"].str.lstrip("0").where(
        df["Código de barras"] != "",
        ("ID:" + df["Identificador interno"]).where(df["Identificador interno"] != "",
                                                    "SKU:" + df["WB SKU"]))
    # Sin código, ID ni SKU (altas desde una tienda): nombre y talla.
    texto = "TXT:" + df["WB Descripcion"].str.upper() + "|" + df["WB Talla"].str.upper()
    return llave.where(llave != "SKU:", texto.where(df["WB Descripcion"] != "", ""))


def limpiar_agregado(texto, marca):
    """Quita de una celda de texto los segmentos que agregó una corrida anterior."""
    return " | ".join(p for p in texto.split(" | ") if not p.startswith(marca)).strip()


AUTOR = "homologar_ciclico_g1522"


def segmento(texto, marca):
    return " | ".join(p for p in texto.split(" | ") if p.startswith(marca))


def aplicar_ciclicos_por_marca(productos):
    """Los Cíclicos por marca (Ciclico_1522_<Marca>.xlsx) son la versión de trabajo: lo que
    se corrigió ahí en los 26 atributos manda sobre el Cíclico integrado. Existencias,
    ubicaciones, notas y fuentes siempre salen del Cíclico integrado."""
    archivos = [f for f in sorted(BASE.glob("Ciclico_1522_*.xlsx")) if f != CICLICO]
    if not archivos:
        return productos
    trabajo = pd.concat([leer_hoja(pd.ExcelFile(f, engine="calamine"), "Productos")
                         [productos.columns] for f in archivos], ignore_index=True)
    trabajo = trabajo[(trabajo[ATRIBUTOS] != "").any(axis=1)]
    # La misma fila se reconoce por código, ID, SKU o nombre y talla, en ese orden, para
    # que un cambio en una de esas llaves no la duplique.
    llaves = {
        "cb": lambda d: d["Código de barras"].str.lstrip("0"),
        "id": lambda d: d["Identificador interno"],
        "sku": lambda d: d["WB SKU"].str.upper(),
        "txt": lambda d: (d["WB Descripcion"].str.upper() + "|" + d["WB Talla"].str.upper())
        .where(d["WB Descripcion"] != "", ""),
    }
    indice = {}
    for nombre, f in llaves.items():
        valores = f(productos)
        indice[nombre] = defaultdict(list)
        for i, v in valores[valores != ""].items():
            indice[nombre][v].append(i)
    columnas = list(productos.columns)
    productos["_nota_previa"] = ""
    productos["_fuente_previa"] = ""
    usadas, nuevas = set(), []
    for _, fila in trabajo.iterrows():
        destino = None
        for nombre, f in llaves.items():
            valor = f(fila.to_frame().T).iloc[0]
            libres = [i for i in indice[nombre].get(valor, []) if i not in usadas] if valor else []
            if libres:
                destino = libres[0]
                break
        if destino is None:
            # Las variantes solo Shopify de Ariat las vuelve a agregar integrar_base_ariat.
            if not fila["Origen del registro"].startswith("Nuevo – Shopify Ariat"):
                nuevas.append(fila[columnas])
        else:
            usadas.add(destino)
            productos.loc[destino, ATRIBUTOS] = fila[ATRIBUTOS].to_numpy()
            # Lo que ya se explicó en una corrida anterior (notas y fuentes web) se conserva.
            productos.loc[destino, "_nota_previa"] = segmento(
                fila["Notas de enriquecimiento / revisión"], MARCA_NOTA)
            productos.loc[destino, "_fuente_previa"] = segmento(fila["Fuentes de consulta"], MARCA_WEB)
    if nuevas:
        productos = pd.concat([productos, pd.DataFrame(nuevas)], ignore_index=True).fillna("")
    print(f"Cíclicos por marca: {len(usadas)} filas actualizadas, {len(nuevas)} nuevas")
    return productos


def seccion(criterios, titulo, salto):
    """Filas de una tabla de Criterios (columnas A–D) hasta la siguiente sección."""
    inicio = criterios.index[criterios[0].str.startswith(titulo)][0] + salto
    filas = []
    for _, r in criterios.loc[inicio:].iterrows():
        if not r[0] or re.match(r"\d+\.\s", r[0]):
            break
        filas.append(r)
    return filas


def leer_criterios(criterios):
    """Lo que la hoja Criterios aporta al proceso.

    · 3 Categorías: familia de talla de cada categoría.
    · 8 Catálogo de tallas: parejas (familia, WB Talla, Talla de EE. UU.).
    · 9 Ubicaciones: bodega (25 / 43) de cada ubicación, columnas G y K.
    · 12 Listas automáticas: valores de los desplegables de Productos, columnas S a AD.
    """
    familias = {r[0]: r[2] for r in seccion(criterios, "3. CATEG", 3)}
    tallas = [(r[0], r[1], r[2]) for r in seccion(criterios, "8. CAT", 3)]
    encabezado = criterios.index[criterios[6] == "Ubicación"][0]
    bodegas = {r[6]: r[10] for _, r in criterios.loc[encabezado + 1:].iterrows()
               if r[6] and r[10]}
    listas = {}
    for j in range(18, 30):
        valores = [v for v in criterios.loc[encabezado:, j] if v]
        listas[valores[0]] = [v for v in valores[1:] if v not in g.TEXTOS_SUSTITUTOS]
    return {"familias": familias, "tallas": tallas, "bodegas": bodegas, "listas": listas,
            "parejas": {(mx, us) for _, mx, us in tallas}}


def leer_ariat():
    xl = pd.ExcelFile(ARIAT, engine="calamine")
    base = leer_hoja(xl, "Productos")
    revision = leer_hoja(xl, "Revisión")
    return base, revision


# --------------------------------------------------------------------------
# Ariat: evidencia de su tienda Shopify tomada de la base homologada
# --------------------------------------------------------------------------

def integrar_base_ariat(df, base, rev_ariat, rev):
    """Suma a Ariat lo que solo vive en Base_Unificada_Ariat_G1522.xlsx.

    Copia los IDs de Shopify, agrega las variantes que solo existen en Shopify y
    guarda para «completar_desde_ariat» los valores que se tomaron de la tienda.
    """
    for c in SHOPIFY_ARIAT:
        df[c] = ""
    empatada = base[base["Origen de la fusión"] != "Solo Shopify"]
    por_id = empatada[empatada["Identificador interno"] != ""].drop_duplicates(
        "Identificador interno").set_index("Identificador interno")
    hay = df["Identificador interno"].isin(por_id.index)
    for c in SHOPIFY_ARIAT:
        df.loc[hay, c] = df.loc[hay, "Identificador interno"].map(por_id[c])

    solo = base[base["Origen de la fusión"] == "Solo Shopify"].copy()
    solo["Origen del registro"] = "Nuevo – Shopify Ariat (sin empate en maestro)"
    solo = solo[[c for c in df.columns if c in solo.columns]]
    inicio = df["_uid"].max() + 1
    solo["_uid"] = range(inicio, inicio + len(solo))
    df = pd.concat([df, solo], ignore_index=True).fillna("")

    # Las decisiones sobre la tienda Shopify Ariat se conservan tal como se tomaron.
    tipos = {"Conflicto maestro-Shopify", "Conflicto Shopify"}
    campos_tienda = {"Código de barras", "WB Composición", "WB Fit"}
    conserva = rev_ariat[rev_ariat["Tipo"].isin(tipos) | (
        (rev_ariat["Tipo"] == "Requiere revisión") & rev_ariat["Campo"].isin(campos_tienda))]
    uid_cb = df[df["Código de barras"] != ""].drop_duplicates("Código de barras") \
        .set_index("Código de barras")["_uid"]
    uid_sku = df[df["WB SKU"] != ""].drop_duplicates("WB SKU").set_index("WB SKU")["_uid"]
    for fila in conserva.to_dict("records"):
        uid = uid_cb.get(fila["Código de barras"], uid_sku.get(fila["WB SKU"], -1))
        rev.add({"_uid": uid, "Código de barras": fila["Código de barras"],
                 "WB SKU": fila["WB SKU"], "Identificador interno": fila["Identificador interno"],
                 "WB Marca": "Ariat"}, fila["Campo"], fila["Valor origen"], fila["Valor final"],
                fila["Tipo"], fila["Incidencia / acción"], fila["Evidencia"])

    completados = rev_ariat[rev_ariat["Tipo"] == "Completado desde Shopify"]
    evidencia = {(r["Identificador interno"], r["Campo"]): r["Evidencia"]
                 for _, r in completados.iterrows() if r["Identificador interno"]}
    return df, (por_id, evidencia)


def completar_desde_ariat(df, ariat, rev):
    """Llena campos vacíos de Ariat con lo que la base Ariat tomó de Shopify."""
    valores, evidencia = ariat
    fila_de = {ident: i for i, ident in df["Identificador interno"].items() if ident}
    for (ident, campo), ev in evidencia.items():
        if ident not in valores.index or ident not in fila_de:
            continue
        nuevo, i = valores.at[ident, campo], fila_de[ident]
        if not df.at[i, campo] and nuevo:
            cambiar(df, i, campo, nuevo, rev, "Completado desde Shopify",
                    "Valor de la tienda Shopify Ariat (Base_Unificada_Ariat_G1522)", ev)


# --------------------------------------------------------------------------
# 2. Marca de las filas sin WB Marca
# --------------------------------------------------------------------------

def prefijo(texto):
    """Prefijo de código del proveedor: «CL/», «AR» en AR2842-220, «WACA» en WACA0213."""
    token = re.split(r"[\s:]+", texto.strip().upper())[0] if texto.strip() else ""
    m = re.match(r"[A-Z]{1,6}/|[A-Z]{2,5}(?=\d)", token)
    return m.group(0) if m else ""


def marca_canonica(texto):
    """La marca de la lista de Criterios escrita igual que «texto» (sin mayúsculas ni acentos)."""
    return {g.sin_acentos(m): m for m in g.MARCAS}.get(g.sin_acentos(texto.strip()), "")


def marca_archivo(r):
    """Marca que decide el libro: Marca principal; si es Multimarca o está vacía, la
    licencia; si la licencia no es una marca, WB Marca (ya completada por nombre o código)."""
    if r["Marca principal"] in g.MARCAS:
        return r["Marca principal"]
    return marca_canonica(r["WB Licencia"]) or r["WB Marca"]


def inferir_marcas(df, rev):
    con_marca = df[df["WB Marca"] != ""]
    por_estilo = con_marca[con_marca["WB N.º de estilo"] != ""].groupby(
        con_marca["WB N.º de estilo"].str.upper())["WB Marca"].agg(set)
    estadistica = defaultdict(Counter)
    for c in ("WB N.º de estilo", "WB SKU", "WB Descripcion"):
        for p, m in zip(con_marca[c].map(prefijo), con_marca["WB Marca"]):
            if p:
                estadistica[p][m] += 1
    por_prefijo = {}
    for p, cuenta in estadistica.items():
        marca, n = cuenta.most_common(1)[0]
        if n >= 5 and n / sum(cuenta.values()) >= 0.97:
            por_prefijo[p] = (marca, n)
    # Nombres de marca de una palabra también al inicio de un código pegado (ARIATMNS…).
    nombres = [(m, g.patron_palabra(m, "" if " " not in m and len(m) >= 5 else r"\b"))
               for m in sorted(g.MARCAS, key=len, reverse=True) if m != "Generico"]
    nombres.append(("CAPSLAB", g.patron_palabra("CAPS")))

    for i in df.index[df["WB Marca"] == ""]:
        r = df.loc[i]
        marca, evidencia = "", ""
        principal = r["Marca principal"]
        licencia = marca_canonica(r["WB Licencia"])
        estilo = r["WB N.º de estilo"].upper()
        if principal in g.MARCAS:
            marca, evidencia = principal, f"Marca principal: {principal}"
        elif licencia:
            marca, evidencia = licencia, f"WB Licencia: {r['WB Licencia']}"
        elif estilo and len(por_estilo.get(estilo, ())) == 1:
            marca = next(iter(por_estilo[estilo]))
            evidencia = f"Mismo WB N.º de estilo que otras filas de {marca}: {r['WB N.º de estilo']}"
        else:
            texto = g.sin_acentos(f"{r['WB Descripcion']} {r['WB Descripción larga']}")
            halladas = {m for m, patron in nombres if patron.search(texto)}
            if len(halladas) == 1:
                marca = halladas.pop()
                evidencia = f"La descripción nombra la marca: {r['WB Descripcion']}"
            else:
                prefijos = {prefijo(r[c]) for c in ("WB N.º de estilo", "WB SKU",
                                                    "WB Descripcion")} - {""}
                marcas = {por_prefijo[p][0] for p in prefijos if p in por_prefijo}
                if len(marcas) == 1:
                    marca = marcas.pop()
                    p = next(p for p in prefijos if p in por_prefijo)
                    evidencia = (f"Prefijo de código «{p}» usado en {por_prefijo[p][1]} filas "
                                 f"de {marca}: {r['WB Descripcion'] or r['WB SKU']}")
        if marca:
            cambiar(df, i, "WB Marca", marca, rev, "Completado por regla",
                    "Marca deducida; confirmar con el proveedor", evidencia)
        else:
            rev.add(r, "WB Marca", "", "", "Requiere revisión",
                    "Sin marca ni evidencia para deducirla; asignar manualmente",
                    f"{r['WB Descripcion']} · {r['Origen del registro']}")


# --------------------------------------------------------------------------
# 3. Normalizar: casos propios del Cíclico
# --------------------------------------------------------------------------

def codigo_de_sku(df, rev):
    """En NetSuite el Nombre/número del hijo es el UPC: un SKU con forma de GTIN válido
    en una variante sin código se usa como su código de barras."""
    for i in df.index[(df["Código de barras"] == "") & ~es_padre(df)]:
        sku = df.at[i, "WB SKU"]
        if sku.isdigit() and 12 <= len(sku) <= 14 and g.digito_control_gtin(sku):
            cambiar(df, i, "Código de barras", sku, rev, "Completado por regla",
                    "SKU con forma de código de barras válido (convención NetSuite: "
                    "Nombre/número del hijo = UPC)", f"WB SKU: {sku}")


KARMAN = re.compile(r"^(\d{2}-\d{3}-\d{4}-\d{4})[ _-]")


def estilo_desde_sku(df, rev):
    """SKU con formato Karman (Stetson, Roper: 01-001-0016-1076 BL-2XL): el estilo es el
    prefijo de 4 bloques, como en la mayoría de las filas que ya traen ambos."""
    con_ambos = df[(df["WB N.º de estilo"] != "") & df["WB SKU"].str.match(KARMAN)]
    prefijo = con_ambos["WB SKU"].str.extract(KARMAN)[0]
    iguales, total = int((con_ambos["WB N.º de estilo"] == prefijo).sum()), len(con_ambos)
    if not total or iguales / total < 0.9:
        return
    for i in df.index[(df["WB N.º de estilo"] == "") & df["WB SKU"].str.match(KARMAN)
                      & ~es_padre(df)]:
        estilo = KARMAN.match(df.at[i, "WB SKU"]).group(1)
        cambiar(df, i, "WB N.º de estilo", estilo, rev, "Completado por regla",
                "Estilo tomado del SKU (formato Karman: los 4 primeros bloques)",
                f"WB SKU {df.at[i, 'WB SKU']}; convención en {iguales} de {total} filas")


def acentuar(palabra):
    correcta = MATERIAL_ESCRITURA.get(palabra.lower())
    if not correcta:
        return palabra
    if palabra.isupper():
        return correcta.upper()
    return correcta.capitalize() if palabra[0].isupper() else correcta


def normalizar_extra(df, rev):
    licencias = df.loc[df["WB Licencia"] != "", "WB Licencia"]
    canon_marca = {g.sin_acentos(m): m for m in g.MARCAS}
    llave = lambda t: re.sub(r"\s+", " ", re.sub(r"\band\b", "&", g.sin_acentos(t))).strip()
    variantes = defaultdict(Counter)
    for valor, n in licencias.value_counts().items():
        variantes[llave(valor)][valor] = n
    canon_licencia = {}
    for k, cuenta in variantes.items():
        canon = canon_marca.get(k) or next((v for v, _ in cuenta.most_common() if "&" in v),
                                           cuenta.most_common(1)[0][0])
        for v in cuenta:
            canon_licencia[v] = canon

    for i in df.index:
        cb = df.at[i, "Código de barras"]
        if cb.isdigit() and len(cb) == 11 and g.digito_control_gtin(cb.zfill(12)):
            cambiar(df, i, "Código de barras", cb.zfill(12), rev, "Corrección aplicada",
                    "Se restituye el cero inicial del UPC (12 dígitos) perdido en Excel")

        for campo in ("Precio de compra", "Precio de venta"):
            precio = df.at[i, campo]
            if re.fullmatch(r"\d+,\d{1,2}", precio):
                cambiar(df, i, campo, precio.replace(",", "."), rev, "Corrección aplicada",
                        "Coma decimal convertida a punto")

        ciclo = df.at[i, "WB Ciclo de vida"]
        if ciclo and ciclo not in g.CICLOS and ciclo.capitalize() in g.CICLOS:
            cambiar(df, i, "WB Ciclo de vida", ciclo.capitalize(), rev, "Corrección aplicada",
                    "Ciclo de vida homologado a la lista de Criterios (Seasonal, Core)")

        if df.at[i, "Unidad"] == "Pieza":
            cambiar(df, i, "Unidad", "PZS", rev, "Corrección aplicada",
                    "Unidad homologada: Pieza → PZS")

        pais = df.at[i, "WB País de origen"]
        if pais and PAISES.get(g.sin_acentos(pais), pais) != pais:
            cambiar(df, i, "WB País de origen", PAISES[g.sin_acentos(pais)], rev,
                    "Corrección aplicada", "País de origen escrito con el nombre en español")

        licencia = df.at[i, "WB Licencia"]
        if licencia and canon_licencia.get(licencia, licencia) != licencia:
            cambiar(df, i, "WB Licencia", canon_licencia[licencia], rev, "Corrección aplicada",
                    "Licencia homologada a una sola escritura")

        composicion = df.at[i, "WB Composición"]
        if composicion:
            nueva = re.sub(r"[A-Za-z]+", lambda m: acentuar(m.group(0)), composicion)
            if nueva != composicion:
                cambiar(df, i, "WB Composición", nueva, rev, "Corrección aplicada",
                        "Nombre del material escrito con su acento")
            porcentajes = [float(p) for p in re.findall(r"(\d+(?:\.\d+)?)\s*%", nueva)]
            if porcentajes and not re.search(r"[:/]", nueva) and \
                    abs(sum(porcentajes) - 100) > 0.01:
                rev.add(df.loc[i], "WB Composición", nueva, nueva, "Requiere revisión",
                        "Los porcentajes no suman 100 % o el texto está truncado; "
                        "confirmar con la etiqueta", f"Suma: {sum(porcentajes):g} %")


# Criterios › 1, pasos 5, 8 y 9: talla de EE. UU., Fit y Silueta.
FIT_CLAVES = [("Relajado", ["relajado", "relaxed"]), ("Slim", ["slim"]),
              ("Clásico", ["clasico", "classic"]), ("Moderno", ["moderno", "modern"]),
              ("Recto", ["recto", "straight"]), ("Bootcut", ["bootcut", "boot cut", "corte bota"]),
              ("Regular", ["regular fit", "corte regular", "fit regular"])]
SILUETA_CLAVES = {
    "Sombreros": [(f"Copa {n}", [n.lower()]) for n in
                  ("Cattleman", "Rancher", "Pinch Front", "Teardrop", "Open Road", "Brick", "Gus")],
    "Calzado": [("Punta Redonda Ancha", ["punta redonda ancha", "broad round toe",
                                         "wide round toe"]),
                ("Punta Cuadrada Ancha", ["punta cuadrada ancha", "broad square toe",
                                          "wide square toe"]),
                ("Punta Cuadrada", ["punta cuadrada", "square toe"]),
                ("Punta Redonda", ["punta redonda", "round toe"]),
                ("Punta Ovalada", ["punta ovalada"]), ("Punta Snip", ["punta snip", "snip toe"]),
                ("Horma Alexia", ["alexia"]), ("Horma Agatha", ["agatha"])],
    "Jeans": [("Bootcut", ["bootcut", "boot cut", "corte bota"]), ("Recto", ["recto", "straight"]),
              ("Skinny", ["skinny"]), ("Tapered", ["tapered"]), ("Trouser", ["trouser"])],
}
CATEGORIA_SILUETA = {"Sombreros": "Sombreros", "Botas": "Calzado", "Zapatos": "Calzado",
                     "Jeans": "Jeans"}


def valores_clave(texto, claves):
    """Valores cuyas palabras clave aparecen; la frase más larga consume a la más corta
    (punta cuadrada ancha no cuenta también como punta cuadrada)."""
    t = " " + g.sin_acentos(texto) + " "
    hallados = []
    frases = sorted(((f, v) for v, fs in claves for f in fs), key=lambda x: -len(x[0]))
    for frase, valor in frases:
        patron = re.compile(r"\b" + re.escape(frase) + r"\b")
        if patron.search(t):
            hallados.append(valor)
            t = patron.sub(" ", t)
    return set(hallados)


def familias_de(categoria, genero, criterios):
    """Familias del catálogo de tallas (Criterios › 8) que corresponden a la categoría
    (Criterios › 3) y, cuando la familia distingue género, al género del producto."""
    texto = criterios["familias"].get(categoria, "")
    salida = set()
    for familia in {f for f, _, _ in criterios["tallas"]}:
        base = re.split(r" – | \(", familia)[0]
        sufijo = familia.split(" – ")[1].split(" ")[0] if " – " in familia else ""
        if base and base in texto and (not sufijo or not genero or sufijo == genero):
            salida.add(familia)
    return salida


def aplicar_criterios(df, rev, criterios):
    for i in df.index[~es_padre(df)]:
        r = df.loc[i]
        categoria, division = r["WB Categoría"], r["WB División"]
        talla, us = r["WB Talla"], r["WB Talla de EE. UU."]

        # Paso 5: la talla de EE. UU. sale de la misma fila del catálogo.
        if categoria and talla and not us:
            familias = familias_de(categoria, r["WB Género"], criterios)
            opciones = {u for f, mx, u in criterios["tallas"] if f in familias and mx == talla}
            if len(opciones) == 1:
                cambiar(df, i, "WB Talla de EE. UU.", opciones.pop(), rev, "Completado por regla",
                        "Talla de EE. UU. de la misma fila del catálogo de tallas (Criterios › 8)",
                        f"{categoria} · WB Talla {talla}")

        texto = f"{r['WB Descripcion']} {r['WB Descripción larga']}"
        # Paso 8: Fit solo en Ropa y Denim, con evidencia en la descripción.
        if not r["WB Fit"] and division in ("Ropa", "Denim"):
            fits = valores_clave(texto, FIT_CLAVES)
            if len(fits) == 1:
                cambiar(df, i, "WB Fit", fits.pop(), rev, "Completado por regla",
                        "Fit identificado en la descripción (Criterios › 7); confirmar",
                        texto[:300])
        # Paso 9: Silueta en Sombreros, Botas, Zapatos y Jeans.
        tipo = CATEGORIA_SILUETA.get(categoria)
        if tipo and not r["WB Silueta"]:
            siluetas = valores_clave(texto, SILUETA_CLAVES[tipo])
            if len(siluetas) == 1:
                cambiar(df, i, "WB Silueta", siluetas.pop(), rev, "Completado por regla",
                        "Silueta identificada en la descripción (Criterios › 7); confirmar",
                        texto[:300])


# --------------------------------------------------------------------------
# 4. Integrar productos iguales
# --------------------------------------------------------------------------

def es_padre(df):
    return df["Origen del registro"].str.contains("padre", na=False)


IDS_NETSUITE = set()  # IDs activos en NetSuite al corte actual; los llena main().


def prioridad(r):
    # NetSuite es la fuente de verdad: su artículo se conserva por encima de cualquier otro.
    puntos = 1000 if r["Identificador interno"] in IDS_NETSUITE else 0
    cb = r["Código de barras"]
    if cb.isdigit() and 12 <= len(cb) <= 14 and g.digito_control_gtin(cb):
        puntos += 100
    puntos += next((p for clave, p in PRIORIDAD_ORIGEN if clave in r["Origen del registro"]), 0)
    return puntos + sum(1 for c in ATRIBUTOS if r[c]) / 100


def compatibles(a, b, campos=("WB Marca", "WB N.º de estilo", "WB Talla", "WB Color")):
    return all(not a[c] or not b[c] or a[c].upper() == b[c].upper() for c in campos)


class Grupos:
    def __init__(self):
        self.padre = {}

    def raiz(self, x):
        while self.padre.get(x, x) != x:
            self.padre[x] = self.padre.get(self.padre[x], self.padre[x])
            x = self.padre[x]
        return x

    def unir(self, a, b):
        self.padre[self.raiz(b)] = self.raiz(a)


def integrar_duplicados(df, rev):
    """Funde en un registro los que representan el mismo producto."""
    padre = es_padre(df)
    grupos = Grupos()
    motivos = {}
    filas = df.to_dict("index")

    def ligar(indices, motivo):
        base = indices[0]
        for j in indices[1:]:
            if compatibles(filas[base], filas[j]):
                grupos.unir(base, j)
                motivos.setdefault(grupos.raiz(base), set()).add(motivo)
            else:
                rev.add(filas[j], "Código de barras" if "código" in motivo else "WB SKU",
                        filas[j]["Código de barras"] or filas[j]["WB SKU"],
                        filas[j]["Código de barras"] or filas[j]["WB SKU"],
                        "Requiere revisión",
                        f"Colisión de captura: {motivo} pero estilo, talla o color distintos; "
                        "corregir el valor equivocado",
                        f"Otra fila: ID {filas[base]['Identificador interno']} · estilo "
                        f"{filas[base]['WB N.º de estilo']} · talla {filas[base]['WB Talla']}")

    variantes = df[~padre]
    cb = variantes["Código de barras"].str.lstrip("0")
    for _, idx in variantes[cb != ""].groupby(cb[cb != ""]).groups.items():
        if len(idx) > 1:
            ligar(list(idx), "mismo código de barras")
    sku = variantes["WB SKU"].str.upper()
    for _, idx in variantes[sku != ""].groupby(sku[sku != ""]).groups.items():
        if len(idx) > 1:
            ligar(sorted(idx, key=lambda i: -prioridad(filas[i])), "mismo WB SKU")

    # Variante sin código igual (marca, estilo, talla y color) a una sola con código.
    llave = (variantes["WB Marca"] + "|" + variantes["WB N.º de estilo"].str.upper() + "|"
             + variantes["WB Talla"].str.upper())
    con_codigo = variantes[(variantes["Código de barras"] != "")
                           & (variantes["WB N.º de estilo"] != "")
                           & (variantes["WB Talla"] != "")]
    candidatas = con_codigo.groupby(llave[con_codigo.index]).groups
    for i in variantes.index[(variantes["Código de barras"] == "")
                             & (variantes["WB N.º de estilo"] != "")
                             & (variantes["WB Talla"] != "")]:
        opciones = [j for j in candidatas.get(llave[i], []) if compatibles(filas[i], filas[j])]
        if len(opciones) == 1:
            grupos.unir(opciones[0], i)
            motivos.setdefault(grupos.raiz(opciones[0]), set()).add(
                "variante sin código igual en estilo/talla/color")

    padres = df[padre & (df["WB N.º de estilo"] != "")]
    llave_padre = padres["WB Marca"] + "|" + padres["WB N.º de estilo"].str.upper()
    for _, idx in padres.groupby(llave_padre).groups.items():
        if len(idx) > 1:
            for j in list(idx)[1:]:
                grupos.unir(idx[0], j)
            motivos.setdefault(grupos.raiz(idx[0]), set()).add("padre repetido del mismo estilo")

    miembros = defaultdict(list)
    for i in df.index:
        miembros[grupos.raiz(i)].append(i)

    df["Códigos de barras alternos"] = ""
    df["Registros integrados"] = ""
    absorbidas, destino, integrados_detalle = [], {}, []
    for raiz, indices in miembros.items():
        if len(indices) == 1:
            continue
        indices.sort(key=lambda i: -prioridad(filas[i]))
        p, otros = indices[0], indices[1:]
        motivo = ", ".join(sorted(motivos.get(grupos.raiz(p), {"duplicado"})))
        alternos, integrados = [], []
        for j in otros:
            o = filas[j]
            destino[o["_uid"]] = df.at[p, "_uid"]
            if o["Código de barras"] and o["Código de barras"] != df.at[p, "Código de barras"] \
                    and o["Código de barras"].lstrip("0") != df.at[p, "Código de barras"].lstrip("0"):
                alternos.append(o["Código de barras"])
                rev.add(df.loc[p], "Código de barras", o["Código de barras"],
                        df.at[p, "Código de barras"], "Requiere revisión",
                        "El mismo producto tiene dos códigos de barras; se conserva el del "
                        "registro principal y el otro queda en «Códigos de barras alternos». "
                        "Confirmar cuál trae la etiqueta física",
                        f"ID {o['Identificador interno'] or '—'} · {o['Origen del registro']}")
            integrados.append(f"ID {o['Identificador interno'] or '—'} · "
                              f"{o['Origen del registro'] or 'sin origen'}"
                              + (f" · código {o['Código de barras']}" if o["Código de barras"] else ""))
            for c in CAMPOS_FUSION + SHOPIFY_ARIAT:
                if c not in df.columns or not o[c]:
                    continue
                actual = df.at[p, c]
                if not actual:
                    cambiar(df, p, c, o[c], rev, "Integrado de duplicado",
                            f"Campo vacío completado con el registro duplicado ({motivo})",
                            f"ID {o['Identificador interno'] or '—'} · {o['Origen del registro']}")
                elif g.sin_acentos(actual) != g.sin_acentos(o[c]) and c in ATRIBUTOS:
                    rev.add(df.loc[p], c, o[c], actual, "Conflicto entre duplicados",
                            "Se conserva el valor del registro principal; validar",
                            f"ID {o['Identificador interno'] or '—'} · {o['Origen del registro']}")
            for c in NUMERICAS + ["_ns_anterior"]:
                a, b = g.numero(df.at[p, c]), g.numero(o[c])
                if b is not None:
                    df.at[p, c] = f"{(a or 0) + b:g}"
            for c in TEXTOS_UNION:
                partes = [x for x in (df.at[p, c], o[c]) if x]
                df.at[p, c] = "; ".join(dict.fromkeys(partes))
        df.at[p, "Códigos de barras alternos"] = " | ".join(dict.fromkeys(alternos))
        df.at[p, "Registros integrados"] = f"{motivo}: " + " | ".join(integrados)
        panel = f"{MARCA_NOTA}: integra {len(otros)} registro(s) iguales por {motivo}"
        if alternos:
            panel += "; códigos de barras alternos: " + ", ".join(dict.fromkeys(alternos))
        df.at[p, "🔍 Panel de Duplicados"] = " | ".join(
            x for x in (limpiar_agregado(df.at[p, "🔍 Panel de Duplicados"], MARCA_NOTA), panel) if x)
        integrados_detalle.append((df.at[p, "_uid"], motivo, [filas[p]] + [filas[j] for j in otros]))
        rev.add(df.loc[p], "Registro", f"{len(indices)} registros", "1 registro",
                "Productos integrados", f"Se fundieron por {motivo}",
                " | ".join(integrados))
        absorbidas += otros
    return df.drop(index=absorbidas).reset_index(drop=True), destino, integrados_detalle


def completar_padres(df, rev):
    """El padre toma de sus variantes la clasificación compartida cuando es unánime."""
    padre = es_padre(df)
    variantes = df[~padre & (df["WB N.º de estilo"] != "")]
    llave = variantes["WB Marca"] + "|" + variantes["WB N.º de estilo"].str.upper()
    for campo in ("WB División", "WB Categoría", "WB Género", "WB Temporada",
                  "WB Ciclo de vida"):
        valores = variantes[variantes[campo] != ""].groupby(llave)[campo].agg(set)
        for i in df.index[padre & (df[campo] == "") & (df["WB N.º de estilo"] != "")]:
            k = f"{df.at[i, 'WB Marca']}|{df.at[i, 'WB N.º de estilo'].upper()}"
            if len(valores.get(k, ())) == 1:
                cambiar(df, i, campo, next(iter(valores[k])), rev, "Completado por regla",
                        "Atributo del modelo tomado de sus variantes (todas coinciden)",
                        f"WB N.º de estilo: {df.at[i, 'WB N.º de estilo']}")


# --------------------------------------------------------------------------
# 5. Validar
# --------------------------------------------------------------------------

def validar(df, parejas_talla):
    df = g.validar(df)
    cb, sku = df["Código de barras"], df["WB SKU"]
    rep_cb = cb.duplicated(keep=False) & (cb != "")
    rep_sku = sku.duplicated(keep=False) & (sku != "")
    talla_cero_valida = {"Vestidos", "Faldas", "Jeans", "Pantalones", "Shorts"}
    netsuite, shopify, avisos_talla = [], [], []
    for i, r in df.iterrows():
        if r["Nivel G1522"] != "Variante":
            netsuite.append("")
            shopify.append("")
            avisos_talla.append("")
            continue
        comun = []
        if not r["Código de barras"]:
            comun.append("V1: falta código de barras")
        elif not 12 <= len(r["Código de barras"]) <= 14:
            comun.append("V1: UPC debe tener 12-14 dígitos")
        if rep_cb[i]:
            comun.append("V2: UPC repetido")
        if rep_sku[i]:
            comun.append("V4: SKU repetido")
        if r["Código de barras"] and not r["WB N.º de estilo"]:
            comun.append("V6: falta estilo")
        ns = list(comun)
        talla = r["WB Talla"]
        if not talla or (talla == "0" and r["WB Categoría"] not in talla_cero_valida):
            ns.append("N2: talla cero o vacía")
        tienda = list(comun)
        venta = g.numero(r["Precio de venta"])
        if venta:
            con_iva = round(venta * 1.16, 2)
            if r["WB Marca"] == "Stetson":
                if con_iva % 50:
                    tienda.append("ST: precio con IVA no es múltiplo de 50")
            elif str(int(con_iva))[-1] != "9":
                tienda.append("WB: precio con IVA no termina en 9")
        netsuite.append(semaforo(ns))
        shopify.append(semaforo(tienda))
        us = r["WB Talla de EE. UU."]
        avisos = []
        if talla and us and (talla, us) not in parejas_talla:
            avisos.append("[Aviso] Pareja de tallas fuera del catálogo de Criterios")
        compra = g.numero(r["Precio de compra"])
        if compra and venta and compra < 20 <= venta / 10:
            # Decisión N5 del flujo NetSuite: confirmar la moneda del costo antes de cargar.
            avisos.append("[Aviso] Precio de compra posiblemente en USD")
        avisos_talla.append("; ".join(avisos))

    df["Validador NetSuite / Odoo"] = netsuite
    df["Validador Shopify"] = shopify
    extra = pd.Series(avisos_talla, index=df.index)
    con_aviso = extra != ""
    df.loc[con_aviso, "Incidencias G1522"] = (df.loc[con_aviso, "Incidencias G1522"]
                                              .where(df.loc[con_aviso, "Incidencias G1522"] == "",
                                                     df.loc[con_aviso, "Incidencias G1522"] + "; ")
                                              + extra[con_aviso])
    df.loc[con_aviso & (df["Estatus G1522"] == "Válido"), "Estatus G1522"] = "Válido con avisos"
    return df


def semaforo(motivos):
    return "✅ PASA" if not motivos else "❌ " + "; ".join(motivos)


# --------------------------------------------------------------------------
# Libros por marca
# --------------------------------------------------------------------------

def grupo_de(marca):
    return GRUPOS.get(marca, marca or "Sin marca")


def slug(texto):
    return re.sub(r"[^A-Za-z0-9]+", "_", g.sin_acentos(texto).title()).strip("_") \
        .replace("Capslab", "CAPSLAB").replace("Reflo", "REFLO")


# --------------------------------------------------------------------------
# Investigación en línea (Investigacion_web_G1522.csv, generado por
# investigar_web_g1522.py a partir de las tiendas y sitios oficiales de cada marca)
# --------------------------------------------------------------------------

SILUETAS_DE = {"Sombreros": g.SILUETA_SOMBRERO, "Botas": g.SILUETA_CALZADO,
               "Zapatos": g.SILUETA_CALZADO, "Jeans": g.SILUETA_JEANS}


def completar_desde_web(df, rev):
    """Llena campos vacíos con lo publicado por la marca o su distribuidor. Solo valores
    ya traducidos a las listas de Criterios; la URL queda en «Fuentes de consulta»."""
    if not WEB.exists():
        return
    web = pd.read_csv(WEB, dtype=str, keep_default_na=False)
    norm = lambda t: re.sub(r"[^A-Z0-9]", "", t.upper())
    indices = {
        "Código de barras": df["Código de barras"].str.lstrip("0"),
        "WB SKU": df["WB SKU"].map(norm),
        "WB N.º de estilo": df["WB N.º de estilo"].map(norm),
    }
    mapas = {k: defaultdict(list) for k in indices}
    for llave, serie in indices.items():
        for i, v in serie.items():
            if v:
                mapas[llave][v].append(i)
    usados = Counter()
    for r in web.itertuples(index=False):
        valor = r.Valor_llave.lstrip("0") if r.Llave == "Código de barras" else norm(r.Valor_llave)
        for i in mapas[r.Llave].get(valor, []):
            if r.Llave == "WB N.º de estilo" and df.at[i, "WB Marca"] and \
                    grupo_de(df.at[i, "WB Marca"]) != grupo_de(r.Marca):
                continue
            if df.at[i, r.Campo] or es_padre(df.loc[[i]]).iloc[0]:
                continue
            # Silueta solo de la lista de su categoría (Criterios › 7): «Rancher» es copa
            # de sombrero, no horma de bota.
            if r.Campo == "WB Silueta" and r.Valor not in SILUETAS_DE.get(
                    df.at[i, "WB Categoría"], set()):
                continue
            cambiar(df, i, r.Campo, r.Valor, rev, "Completado por investigación en línea",
                    f"Dato publicado por la marca o su distribuidor ({r.Llave}); confirmar",
                    f"{r.Fuente} · {r.Evidencia}")
            fuente = f"{MARCA_WEB} {r.Fecha}: {r.Fuente}"
            actual = df.at[i, "Fuentes de consulta"]
            if fuente not in actual:
                df.at[i, "Fuentes de consulta"] = f"{actual} | {fuente}" if actual else fuente
            usados[r.Campo] += 1
    print("Investigación en línea aplicada:", dict(usados))


# --------------------------------------------------------------------------
# Cíclicos por marca con la estructura de Ciclico_1522_Stetson.xlsx
# --------------------------------------------------------------------------

LISTA = {"División": g.DIVISIONES, "Género": g.GENEROS, "Temporada": g.TEMPORADAS,
         "Ciclo de vida": g.CICLOS}


def detalle_discrepancias(df):
    """Las fórmulas AA y AB de la plantilla, calculadas en Python (libros grandes)."""
    cb, sku = df["Código de barras"], df["WB SKU"]
    rep_cb = cb.duplicated(keep=False) & (cb != "")
    rep_sku = sku.duplicated(keep=False) & (sku != "")
    detalle, estatus = [], []
    for i, r in df.iterrows():
        if not (r["Código de barras"] or r["Identificador interno"] or r["WB SKU"]):
            detalle.append(""); estatus.append(""); continue
        d = []
        if not r["Código de barras"].strip() or "FALTA" in r["Código de barras"].upper():
            d.append("UPC faltante/inválido")
        if rep_cb[i]:
            d.append("Código de barras duplicado")
        if rep_sku[i]:
            d.append("SKU duplicado")
        division, categoria = r["WB División"], r["WB Categoría"]
        d.append("[Aviso] División no capturada ()" if not division else
                 "" if division in g.DIVISIONES else "División no válida")
        if not categoria:
            d.append("[Aviso] Categoría no capturada ()")
        elif categoria not in g.CATEGORIA_DIVISION:
            d.append("Categoría no válida")
        elif division and division != g.CATEGORIA_DIVISION[categoria]:
            d.append("División no coincide con Categoría")
        for campo, nombre, valida, genero in (
                ("WB Género", "Género", g.GENEROS, "o"), ("WB Temporada", "Temporada", g.TEMPORADAS, "a"),
                ("WB Ciclo de vida", "Ciclo de vida", g.CICLOS, "o")):
            if not r[campo]:
                d.append(f"[Aviso] {nombre} no capturad{genero} ()")
            elif r[campo] not in valida:
                d.append(f"{nombre} no válid{genero}")
        d = [x for x in d if x]
        texto = "; ".join(d)
        detalle.append(texto)
        errores = sum(1 for x in d if not x.startswith("[Aviso]"))
        estatus.append("Válido" if not d or not errores else "Con Incidencias")
    return estatus, detalle


def nota_homologacion(uid, cambios, r):
    """Resumen de lo hecho en la fila para «Notas de enriquecimiento / revisión». Conserva
    los cambios explicados en corridas anteriores; la validación se recalcula siempre."""
    previa = re.sub(rf"^{MARCA_NOTA} [\d/]+: ", "", r.get("_nota_previa", "") or "")
    partes = [x for x in previa.split(" · ") if x and not x.startswith(("Validación G1522",
                                                                       "Validador "))]
    for c in cambios.get(uid, []):
        if c["Tipo"] == "Productos integrados":
            continue
        if c["Valor final"] != c["Valor origen"]:
            partes.append(f"{c['Campo']}: «{c['Valor origen']}» → «{c['Valor final']}» ({c['Tipo']})")
        else:
            partes.append(f"{c['Campo']}: {c['Incidencia / acción']} ({c['Tipo']})")
    errores = [e for e in r["Incidencias G1522"].split("; ")
               if e and (not e.startswith("[Aviso]") or "USD" in e or "Pareja" in e)]
    if errores:
        partes.append("Validación G1522: " + "; ".join(errores))
    for col in ("Validador NetSuite / Odoo", "Validador Shopify"):
        if r[col].startswith("❌"):
            partes.append(f"{col}: {r[col][2:]}")
    if not partes:
        return ""
    return f"{MARCA_NOTA} {FECHA}: " + " · ".join(dict.fromkeys(partes))


NUMERO = re.compile(r"-?\d+(\.\d+)?$")


def a_celda(columna, valor):
    if valor == "":
        return None
    if columna in COLUMNAS_NUMERICAS and NUMERO.match(valor):
        return float(valor)
    return valor


COLUMNAS_NUMERICAS = {"Precio de compra", "Precio de venta", "Stock Sistema NetSuite",
                      "Disponible NetSuite", "Precio compra vigente", "Precio venta vigente",
                      "Stock Shopify Stetson México", "Stock Shopify Western Brothers",
                      "Stock Odoo Universal Unique Brands"}


def ajustar_rangos(formula, filas_productos, filas_stock):
    """Lleva los rangos fijos de la plantilla al tamaño del libro de la marca."""
    formula = re.sub(r"(Productos!\$?[A-Z]{1,2}\$?2:\$?[A-Z]{1,2}\$?)(\d+)",
                     lambda m: f"{m.group(1)}{filas_productos}", formula)
    formula = re.sub(r"('Stock por Ubicación'!\$A\$2:\$J\$)(\d+)",
                     lambda m: f"{m.group(1)}{max(filas_stock, 2)}", formula)
    return formula


def limpiar_desde(ws, fila, conservar=()):
    """Borra las celdas desde «fila» (valores y formato), salvo las indicadas."""
    for (r, c) in [k for k in ws._cells if k[0] >= fila]:
        if ws.cell(r, c).coordinate not in conservar:
            del ws._cells[(r, c)]
    for r in [r for r in ws.row_dimensions if r >= fila]:
        del ws.row_dimensions[r]


def escribir_ciclico(ruta, plantilla, grupo, productos, columnas, stock, escaneos,
                     duplicados, vivas, inventario=None):
    import openpyxl
    from openpyxl.formatting.formatting import ConditionalFormattingList
    wb = openpyxl.load_workbook(io.BytesIO(plantilla))
    for nombre in wb.sheetnames:  # Solo las 12 hojas del Cíclico; las de inventario se rehacen.
        if nombre not in HOJAS_CICLICO:
            del wb[nombre]
    n = len(productos) + 1
    m = len(stock) + 1

    # Productos: mismas 44 columnas, una fila por variante.
    ws = wb["Productos"]
    estilo = {c: ws.cell(2, c)._style for c in range(1, len(columnas) + 1)}
    form_aa, form_ab = ws["AA2"].value, ws["AB2"].value
    limpiar_desde(ws, 2)
    for k, fila in enumerate(productos[columnas].itertuples(index=False), start=2):
        for j, (col, valor) in enumerate(zip(columnas, fila), start=1):
            celda = ws.cell(k, j, a_celda(col, valor))
            celda._style = estilo[j]
        if vivas:
            ws.cell(k, 27).value = re.sub(r"(?<![$A-Z])([A-Z]{1,2})2\b", rf"\g<1>{k}", form_aa)
            ws.cell(k, 28).value = re.sub(r"(?<![$A-Z])([A-Z]{1,2})2\b", rf"\g<1>{k}",
                                          re.sub(r"(\$[AE]\$2:\$[AE]\$)\d+", rf"\g<1>{n}", form_ab))
    validaciones = []
    for dv in ws.data_validations.dataValidation:
        col = re.match(r"[A-Z]+", str(dv.sqref).split()[0]).group(0)
        if col in [re.match(r"[A-Z]+", str(v.sqref)).group(0) for v in validaciones]:
            continue
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(f"{col}2:{col}{max(n, 2)}")
        validaciones.append(dv)
    ws.data_validations.dataValidation = validaciones
    reglas = [(str(r.sqref), r.rules) for r in ws.conditional_formatting]
    ws.conditional_formatting = ConditionalFormattingList()
    for rango, rs in reglas:
        col = re.match(r"[A-Z]+", rango).group(0)
        for regla in rs:
            ws.conditional_formatting.add(f"{col}2:{col}{max(n, 2)}", regla)
    ws.auto_filter.ref = f"A1:AR{max(n, 2)}"

    # Productos_Resumen: solo las fórmulas ancla; Sheets calcula el resto.
    ws = wb["Productos_Resumen"]
    anclas = {c: ws[c].value for c in ("A2", "Q2", "R2")}
    limpiar_desde(ws, 2)
    for c, f in anclas.items():
        texto = f.text if hasattr(f, "text") else f
        texto = ajustar_rangos(texto, n, m)
        texto = re.sub(r"(?<![!$A-Z])([A-DQ])2:\1\d+", rf"\g<1>2:\g<1>{n}", texto)
        texto = re.sub(r"ARRAY_CONSTRAIN\((.*), \d+, 16\)", rf"ARRAY_CONSTRAIN(\1, {max(n - 1, 1)}, 16)",
                       texto, flags=re.S)
        ws[c] = texto

    # Stock por Ubicación: renglones de la marca; las columnas derivadas por fórmula.
    ws = wb["Stock por Ubicación"]
    estilo = {c: ws.cell(2, c)._style for c in range(1, 11)}
    form = {c: (ws.cell(2, c).value.text if hasattr(ws.cell(2, c).value, "text")
                else ws.cell(2, c).value) for c in (2, 4, 5, 6, 8, 9)}
    limpiar_desde(ws, 2)
    fila_de = dict(zip(productos["_uid"], range(2, n + 1)))
    for k, r in enumerate(stock.itertuples(index=False), start=2):
        valores = {1: r[0], 3: r[2], 7: r[6], 10: r[9]}
        for c in range(1, 11):
            if c in valores:
                v = valores[c]
                ws.cell(k, c, float(v) if c == 7 and NUMERO.match(str(v)) else (v or None))
            elif c == 9 and not vivas:
                ws.cell(k, c, fila_de.get(r.uid))
            else:
                ws.cell(k, c, ajustar_rangos(re.sub(r"(?<![$A-Z])([A-Z])2\b", rf"\g<1>{k}",
                                                    form[c]), n, m))
            ws.cell(k, c)._style = estilo[c]
    ws.auto_filter.ref = f"A1:J{max(m, 2)}"

    # Bodegas: el filtro por bodega es una fórmula de Sheets sobre Stock por Ubicación.
    for nombre in ("Bodega 25", "Bodega 43"):
        ws = wb[nombre]
        ancla = ws["A6"].value
        ancla = ancla.text if hasattr(ancla, "text") else ancla
        totales = {c: ws[c].value for c in ("E3", "G3")}
        limpiar_desde(ws, 6)
        ws["A6"] = ajustar_rangos(ancla, n, m)
        fin = 6 + m + 500
        for c, f in totales.items():
            ws[c] = re.sub(r"I6:I\d+", f"I6:I{fin}", f)

    # Escaneos (tabla Table_2) e Historial: solo los folios con piezas de la marca.
    ws = wb["Escaneos"]
    tabla = ws.tables["Table_2"]
    estilo = {c: ws.cell(2, c)._style for c in range(1, 14)}
    limpiar_desde(ws, 2)
    for k, r in enumerate(escaneos.itertuples(index=False), start=2):
        for c, v in enumerate(r[:13], start=1):
            ws.cell(k, c, float(v) if c == 10 and NUMERO.match(str(v)) else (v or None))
            ws.cell(k, c)._style = estilo[c]
    tabla.ref = f"A1:M{max(len(escaneos) + 1, 2)}"
    ws = wb["Historial de Escaneos"]
    folios = set(escaneos["Folio de Escaneo"])
    filas = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, 3).value in folios:
            filas.append([ws.cell(r, c).value for c in range(1, 10)])
    estilo = {c: ws.cell(4, c)._style for c in range(1, 10)}
    limpiar_desde(ws, 4)
    for k, valores in enumerate(filas, start=4):
        for c, v in enumerate(valores, start=1):
            if isinstance(v, str) and v.startswith("="):
                v = re.sub(r"(?<![$A-Z])C\d+\b", f"C{k}", v)
            ws.cell(k, c, v)._style = estilo[c]

    # Revisión Duplicados: grupos del Cíclico y registros iguales integrados en uno.
    ws = wb["Revisión Duplicados"]
    estilo = {c: ws.cell(2, c)._style for c in range(1, 48)}
    limpiar_desde(ws, 2)
    for k, fila in enumerate(duplicados, start=2):
        for c, v in enumerate(fila, start=1):
            col = (["Grupo", "Fuente", "Motivo"] + columnas)[c - 1]
            ws.cell(k, c, a_celda(col, str(v)) if c > 3 else v)._style = estilo[c]
    ws.auto_filter.ref = f"A1:AU{max(len(duplicados) + 1, 2)}"

    if inventario:
        inv.escribir_inventario(wb, *inventario, grupo)
    wb.properties.creator = AUTOR
    wb.save(ruta)


def main():
    rev = Revision()
    plantilla = PLANTILLA.read_bytes()  # Se lee antes de reescribir el Cíclico de Stetson.
    productos, hojas, criterios = leer_ciclico()
    for c in ("_nota_previa", "_fuente_previa"):
        if c not in productos.columns:
            productos[c] = ""
    columnas = [c for c in productos.columns if not c.startswith("_")]
    print(f"Cíclico: {len(productos)} filas con datos")
    productos["_uid"] = range(len(productos))

    # Inventario al corte más reciente (NetSuite 06/10) y artículos nuevos de NetSuite.
    catalogo_ns, detalle_ns, ubicaciones_ns = inv.leer_netsuite()
    IDS_NETSUITE.update(set(catalogo_ns["Identificador interno"]) - {""})
    # Existencia del corte anterior por artículo: la suma de su desglose por ubicación en el
    # Cíclico (la columna «Stock Sistema NetSuite» de Productos no cuadra con ese desglose en
    # Happy Socks, CAPSLAB y REFLO).
    stock_28 = hojas["Stock por Ubicación"]
    # Renglones con código e ID en «0»: el ID real viene en el texto de origen.
    sin_id = stock_28["Identificador interno"].isin(["", "0"])
    stock_28.loc[sin_id, "Identificador interno"] = stock_28.loc[sin_id, "Origen en Productos"] \
        .str.extract(r"ID interno (\d+)")[0].fillna("")
    stock_28.loc[stock_28["Código de barras"] == "0", "Código de barras"] = ""
    ns_anterior = stock_28[stock_28["Sistema"] == "NetSuite"]
    anterior_por_id = inv.numero(ns_anterior["Stock en ubicación (Físico)"]).groupby(
        ns_anterior["Identificador interno"]).sum().to_dict()
    shopify_ariat = inv.leer_shopify_ariat()
    productos = inv.actualizar_existencias(productos, catalogo_ns, ATRIBUTOS, cambiar, rev,
                                           len(productos), anterior_por_id)

    base_ariat, rev_ariat = leer_ariat()
    productos, ariat = integrar_base_ariat(productos, base_ariat, rev_ariat, rev)
    print(f"Con variantes solo Shopify Ariat: {len(productos)}")

    inferir_marcas(productos, rev)
    codigo_de_sku(productos, rev)
    estilo_desde_sku(productos, rev)
    normalizar_extra(productos, rev)
    completar_desde_web(productos, rev)
    g.normalizar(productos, rev)
    completar_desde_ariat(productos, ariat, rev)  # La ficha de Shopify es mejor evidencia.
    aplicar_criterios(productos, rev, criterios)
    # Segunda pasada: datos web que requieren la categoría ya asignada (Silueta) y la
    # División o Unidad de las categorías que llegaron de la web.
    completar_desde_web(productos, rev)
    g.normalizar(productos, rev)
    llaves_originales = productos[["_uid", "Identificador interno", "Código de barras", "WB SKU",
                                   "Shopify ID variante"]].copy()
    productos, destino, integrados = integrar_duplicados(productos, rev)
    inv.recalcular_netsuite(productos, llaves_originales, destino, catalogo_ns, anterior_por_id)
    print(f"Tras integrar productos iguales: {len(productos)}")
    productos = validar(productos, criterios["parejas"])

    revision = rev.tabla()
    revision["_uid"] = revision["_uid"].map(lambda u: destino.get(u, u))
    cambios = defaultdict(list)
    for c in revision.to_dict("records"):
        cambios[c["_uid"]].append(c)
    productos["Notas de enriquecimiento / revisión"] = [
        " | ".join(x for x in (limpiar_agregado(r["Notas de enriquecimiento / revisión"], MARCA_NOTA),
                               nota_homologacion(r["_uid"], cambios, r)) if x)
        for _, r in productos.iterrows()]

    productos["Fuentes de consulta"] = [
        " | ".join(dict.fromkeys(x for x in f"{actual} | {previa}".split(" | ") if x))
        for actual, previa in zip(productos["Fuentes de consulta"], productos["_fuente_previa"])]
    productos["Libro"] = productos.apply(marca_archivo, axis=1).map(grupo_de)
    variantes = productos[productos["Nivel G1522"] == "Variante"].copy()

    # Inventario, stock por ubicación y escaneos se ligan al registro integrado por ID de
    # NetSuite, código de barras (también los alternos), SKU o ID de variante de Shopify;
    # las llaves de los registros absorbidos apuntan al que los integró.
    final = llaves_originales["_uid"].map(lambda u: destino.get(u, u))
    vivos = set(variantes["_uid"])

    def mapa(serie, transformar=lambda x: x):
        m = {}
        for valor, uid in zip(serie.map(transformar), final):
            if valor and uid in vivos:
                m.setdefault(valor, uid)
        return m

    uid_id = mapa(llaves_originales["Identificador interno"])
    llave_cb = mapa(llaves_originales["Código de barras"], lambda x: x.lstrip("0"))
    for uid, alternos in variantes[["_uid", "Códigos de barras alternos"]].itertuples(index=False):
        for c in alternos.split(" | "):
            if c:
                llave_cb.setdefault(c.lstrip("0"), uid)
    uid_sku = mapa(llaves_originales["WB SKU"])
    uid_variante = mapa(llaves_originales["Shopify ID variante"])

    def ligar(tabla, cb, ident, sku, variante=None):
        u = tabla[cb].str.lstrip("0").map(llave_cb)
        if variante:
            u = tabla[variante].map(uid_variante).fillna(u)
        u = u.fillna(tabla[ident].map(uid_id)) if ident else u
        return u.fillna(tabla[sku].map(uid_sku))

    libro_uid = variantes.set_index("_uid")["Libro"]
    detalle_ns["uid"] = detalle_ns["Identificador interno"].map(uid_id)
    detalle_ns["uid"] = detalle_ns["uid"].fillna(ligar(detalle_ns, "Código de barras", None, "WB SKU"))
    shopify_ariat["uid"] = ligar(shopify_ariat, "Código de barras", None, "WB SKU", "Variant ID")
    stock_ciclico = hojas["Stock por Ubicación"].copy()
    stock_ciclico["uid"] = ligar(stock_ciclico, "Código de barras", "Identificador interno", "WB SKU")
    anterior_ubicacion = stock_ciclico[stock_ciclico["Sistema"] == "NetSuite"].rename(
        columns={"Ubicación (NetSuite / Plataforma)": "Ubicación"})
    anterior_ubicacion["Stock"] = inv.numero(anterior_ubicacion["Stock en ubicación (Físico)"])
    plataformas = stock_ciclico[stock_ciclico["Sistema"] != "NetSuite"].copy()
    plataformas["Stock en ubicación (Físico)"] = inv.numero(plataformas["Stock en ubicación (Físico)"])
    inv.marcar_shopify_ariat(productos, shopify_ariat)
    variantes = productos[productos["Nivel G1522"] == "Variante"].copy()
    anterior_total = pd.Series(inv.numero(variantes["_ns_anterior"]).to_numpy(), index=variantes["_uid"])
    print(f"Inventario ligado: NetSuite {detalle_ns['uid'].notna().mean():.1%} de renglones, "
          f"Shopify Ariat {shopify_ariat['uid'].notna().mean():.1%}")

    stock = inv.stock_por_ubicacion(detalle_ns, stock_ciclico, shopify_ariat)
    stock = stock[list(hojas["Stock por Ubicación"].columns) + ["uid"]]
    stock["Libro"] = stock["uid"].map(libro_uid).fillna(stock["WB Marca"].map(grupo_de))
    ubic_info = {}
    for _, u in ubicaciones_ns.iterrows():
        tipo = criterios["bodegas"].get(u["Ubicación corta"]) or u["Tipo"]
        ubic_info[u["Ubicación corta"]] = {"tipo": tipo, "activa": u["Activa"]}
    escaneos = hojas["Escaneos"].copy()
    escaneos["Piezas Contadas"] = inv.numero(escaneos["Piezas Contadas"])
    escaneos["uid"] = escaneos["UPC CODE"].str.lstrip("0").map(llave_cb)
    escaneos["Libro"] = escaneos["uid"].map(libro_uid).fillna(escaneos["Marca"].map(grupo_de))
    dup = hojas["Revisión Duplicados"]
    # Un grupo de duplicados va completo a cada archivo de las marcas que toca.
    libros_grupo = dup.assign(L=dup["WB Marca"].map(grupo_de)).groupby("Grupo")["L"].agg(set)

    if not sys.argv[1:]:
        for viejo in BASE.glob("Ciclico_1522_*.xlsx"):
            if viejo not in (CICLICO, PLANTILLA):
                viejo.unlink()
    resumen = []
    solo = sys.argv[1:]  # Opcional: generar solo estas marcas, p. ej. «Stetson Roper».
    for grupo, filas in variantes.groupby("Libro", sort=False):
        if solo and slug(grupo) not in solo and grupo not in solo:
            continue
        ruta = BASE / f"Ciclico_1522_{slug(grupo)}.xlsx"
        vivas = len(filas) <= LIMITE_FORMULAS
        filas = filas.copy()
        if not vivas:
            filas["Estatus de Validación"], filas["Detalle de Discrepancias"] = \
                detalle_discrepancias(filas.reset_index(drop=True))
        uids = set(filas["_uid"])
        filas_dup = [[r["Grupo"], r["Fuente"], r["Motivo"]] + [r[c] for c in columnas]
                     for _, r in dup[dup["Grupo"].map(lambda k: grupo in libros_grupo[k])].iterrows()]
        siguiente = max([int(float(x[0])) for x in filas_dup if str(x[0]).replace(".0", "").isdigit()]
                        or [0]) + 1
        for uid, motivo, registros in integrados:
            if uid not in uids:
                continue
            for k, reg in enumerate(registros):
                filas_dup.append([siguiente, "Productos — registro conservado" if k == 0 else
                                  "Productos — registro integrado en el anterior",
                                  f"{MARCA_NOTA}: mismo producto ({motivo})"]
                                 + [reg.get(c, "") for c in columnas])
            siguiente += 1
        stock_marca = stock[stock["Libro"] == grupo][
            list(hojas["Stock por Ubicación"].columns) + ["uid"]]
        escaneos_marca = escaneos[escaneos["Libro"] == grupo]
        inventario, ubicaciones, tiendas, pedido = inv.tabla_inventario(
            filas, detalle_ns, anterior_total, plataformas, shopify_ariat, escaneos)
        movimientos = inv.movimientos_por_ubicacion(detalle_ns, anterior_ubicacion, uids)
        escribir_ciclico(ruta, plantilla, grupo, filas, columnas, stock_marca,
                         escaneos_marca, filas_dup, vivas,
                         (inventario, ubicaciones, tiendas, pedido, ubic_info, movimientos))
        estatus = filas["Estatus G1522"].value_counts()
        resumen.append({"Archivo": ruta.name, "Variantes": len(filas),
                        "Fórmulas vivas": "Sí" if vivas else "Validación calculada",
                        "Válido": int(estatus.get("Válido", 0)),
                        "Válido con avisos": int(estatus.get("Válido con avisos", 0)),
                        "Con incidencias": int(estatus.get("Con incidencias", 0)),
                        "Productos integrados": int((filas["Registros integrados"] != "").sum()),
                        "Stock por ubicación": len(stock_marca), "Escaneos": len(escaneos_marca),
                        "Artículos en Inventario": len(inventario),
                        "En mano NetSuite": int(inventario["En mano NetSuite"].sum()),
                        "MB": round(ruta.stat().st_size / 1e6, 2)})
        print(resumen[-1], flush=True)
    print(pd.DataFrame(resumen).to_string(index=False))


if __name__ == "__main__":
    main()
