"""Inventario actual por marca para los Cíclicos G1522.

Homologa en un solo lugar (el Cíclico de cada marca) el inventario de todas las fuentes:

  · NetSuite 06/10/2026 (Catalogo_Estandar_Inventario_NetSuite_2026-10-06.xlsx): existencia
    oficial por ubicación: en mano, disponible, comprometido, en pedido, en tránsito y
    pendiente por surtir. Es el inventario del sistema.
  · NetSuite 28/09/2026 (corte del Cíclico integrado): sirve para medir el movimiento neto.
  · Shopify Ariat 05/10/2026 (Products.csv): existencia por tienda Ariat.
  · Shopify Stetson México, Shopify Western Brothers y Odoo 01/10/2026 (Cíclico integrado).
  · Escaneos del conteo cíclico (piezas contadas en rack).

homologar_ciclico_g1522.py usa estas funciones para actualizar las columnas de existencia de
Productos, rehacer Stock por Ubicación y agregar las hojas «Inventario» y «KPIs Inventario».
"""

import re
from pathlib import Path

import pandas as pd

BASE = Path(__file__).resolve().parent
INVENTARIO_NS = BASE / "Catalogo_Estandar_Inventario_NetSuite_2026-10-06.xlsx"
SHOPIFY_ARIAT = BASE / "Products.csv"
CORTE_NS = "06/10/2026"
CORTE_ANTERIOR = "28/09/2026"
CORTE_SHOPIFY_ARIAT = "05/10/2026"
CORTE_PLATAFORMAS = "01/10/2026"

CANTIDADES = ["En mano", "Disponible", "Comprometido", "En pedido", "En tránsito",
              "Pendiente por surtir"]
# Columnas de existencia de Productos que se toman de NetSuite al corte más reciente.
STOCK_PRODUCTOS = ["Stock Sistema NetSuite", "Ubicaciones con stock (NetSuite)",
                   "Disponible NetSuite", "¿Alguna ubicación en negativo?",
                   "Ubicaciones con existencia negativa", "Precio compra vigente",
                   "Precio venta vigente"]
# Tiendas Ariat que existen en NetSuite y en Shopify Ariat (pareja por nombre).
TIENDAS_ARIAT = {"ARIAT CHIHUAHUA": "Ariat Chihuahua", "ARIAT OUTLET GDL": "Ariat Outlet GDL",
                 "ARIAT PLAZA SAN LUIS": "Ariat Plaza San Luis",
                 "ARIAT SAN NICOLAS NL": "Ariat San Nicolas de los Garza",
                 "Nogales Store": "Ariat Ecuestre Nogales", "Plaza del Sol Store": "Plaza del sol"}
PLATAFORMAS = ["Shopify Stetson México", "Shopify Western Brothers", "Odoo Universal Unique Brands"]


def ubicacion_corta(nombre):
    """«CAPSLAB Warehouse : CL Muestras (inactiva)» → «CL Muestras», como en Criterios › 9."""
    return re.sub(r"\s*\(inactiva\)$", "", nombre.split(" : ")[-1]).strip()


def numero(serie):
    return pd.to_numeric(serie, errors="coerce").fillna(0)


# --------------------------------------------------------------------------
# Lectura
# --------------------------------------------------------------------------

def leer_netsuite():
    xl = pd.ExcelFile(INVENTARIO_NS, engine="calamine")
    leer = lambda h: pd.read_excel(xl, sheet_name=h, dtype=str).fillna("").apply(
        lambda s: s.str.strip())
    catalogo = leer("Catálogo estándar")
    detalle = leer("Inventario detalle")
    for c in CANTIDADES:
        detalle[c] = numero(detalle[c])
    detalle["Ubicación corta"] = detalle["Ubicación"].map(ubicacion_corta)
    ubicaciones = leer("Ubicaciones")
    ubicaciones = ubicaciones[ubicaciones["ID interno"].str.fullmatch(r"\d+")]
    ubicaciones["Ubicación corta"] = ubicaciones["Ubicación (etiqueta en este libro)"].map(
        ubicacion_corta)
    return catalogo, detalle, ubicaciones


def leer_shopify_ariat():
    """Existencia por tienda de la exportación Matrixify de Shopify Ariat (renglón largo)."""
    sp = pd.read_csv(SHOPIFY_ARIAT, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    sp["Title"] = sp["Title"].replace("", None)
    sp["Title"] = sp.groupby("ID")["Title"].transform("first").fillna("")
    var = sp[sp["Variant ID"] != ""]
    filas = []
    for columna in var.columns:
        m = re.fullmatch(r"Inventory On Hand: (.+)", columna)
        if not m:
            continue
        tienda = m.group(1)
        datos = {"En mano": numero(var[columna]),
                 "Disponible": numero(var.get(f"Inventory Available: {tienda}", 0)),
                 "Comprometido": numero(var.get(f"Inventory Committed: {tienda}", 0)),
                 "Por llegar": numero(var.get(f"Inventory Incoming: {tienda}", 0))}
        hay = (datos["En mano"] != 0) | (datos["Por llegar"] != 0) | (datos["Comprometido"] != 0)
        tabla = pd.DataFrame({
            "Tienda": tienda, "Ubicación": f"Shopify Ariat México › {tienda}",
            "Variant ID": var["Variant ID"], "Código de barras": var["Variant Barcodes"],
            "WB SKU": var["Variant SKU"], "Descripción": var["Title"], **datos})[hay]
        filas.append(tabla)
    return pd.concat(filas, ignore_index=True)


# --------------------------------------------------------------------------
# Productos al corte más reciente
# --------------------------------------------------------------------------

def actualizar_existencias(productos, catalogo, atributos, cambiar, rev, siguiente_uid,
                           anterior_por_id):
    """Existencias de Productos al corte NetSuite 06/10; agrega los artículos nuevos.

    La existencia 28/09 (suma por ubicación del Cíclico) queda en «_ns_anterior» para medir
    el movimiento. Un campo vacío
    del Cíclico se completa con lo capturado en NetSuite; nunca se sobrescribe un valor.
    """
    productos["_ns_anterior"] = productos["Identificador interno"].map(anterior_por_id).fillna(0).where(
        ~productos["Origen del registro"].str.startswith(f"Nuevo – NetSuite {CORTE_NS}"), 0)
    ns = catalogo.drop_duplicates("Identificador interno").set_index("Identificador interno")
    hay = (productos["Identificador interno"] != "") & productos["Identificador interno"].isin(ns.index)
    ids = productos.loc[hay, "Identificador interno"]
    for c in STOCK_PRODUCTOS:
        productos.loc[hay, c] = ids.map(ns[c]).to_numpy()
    for c in atributos:
        if c not in ns.columns:
            continue
        nuevo = ids.map(ns[c])
        for i in nuevo.index[(productos.loc[hay, c] == "") & (nuevo != "")]:
            cambiar(productos, i, c, nuevo[i], rev, "Completado desde NetSuite",
                    f"Campo vacío capturado en NetSuite al {CORTE_NS}",
                    f"NetSuite ID interno {productos.at[i, 'Identificador interno']}")
    # Artículos de NetSuite que ya no están activos: sin existencia en el corte actual.
    de_netsuite = productos["Origen del registro"].str.contains("NetSuite|Catálogo existente|Hoja a Revisar")
    fuera = de_netsuite & (productos["Identificador interno"] != "") & ~hay
    for c in ("Stock Sistema NetSuite", "Disponible NetSuite"):
        productos.loc[fuera, c] = "0"
    for c in ("Ubicaciones con stock (NetSuite)", "Ubicaciones con existencia negativa"):
        productos.loc[fuera, c] = ""
    productos.loc[fuera, "¿Alguna ubicación en negativo?"] = "No"

    conocidos_cb = set(productos["Código de barras"].str.lstrip("0")) - {""}
    nuevos = catalogo[~catalogo["Identificador interno"].isin(productos["Identificador interno"])
                      & ~catalogo["Código de barras"].str.lstrip("0").isin(conocidos_cb)].copy()
    nuevos["Origen del registro"] = f"Nuevo – NetSuite {CORTE_NS}"
    nuevos = nuevos.reindex(columns=productos.columns, fill_value="")
    nuevos["_ns_anterior"] = 0
    nuevos["_uid"] = range(siguiente_uid, siguiente_uid + len(nuevos))
    print(f"NetSuite {CORTE_NS}: {int(hay.sum())} filas actualizadas, {int(fuera.sum())} ya no "
          f"activas (existencia 0), {len(nuevos)} artículos nuevos")
    return pd.concat([productos, nuevos], ignore_index=True).fillna("")


def recalcular_netsuite(productos, llaves, destino, catalogo, anterior_por_id):
    """Existencia NetSuite de cada registro = suma de los artículos de NetSuite (IDs) que
    contiene tras integrar duplicados. Un registro absorbido sin ID propio no suma: su
    existencia era copia de la del artículo que lo absorbió."""
    ns = catalogo.drop_duplicates("Identificador interno").set_index("Identificador interno")
    ids = llaves.assign(final=llaves["_uid"].map(lambda u: destino.get(u, u)))
    # Corte anterior: todos los IDs que forman el registro, estén o no activos hoy.
    todos = ids[ids["Identificador interno"] != ""].drop_duplicates("Identificador interno")
    anterior = todos["Identificador interno"].map(anterior_por_id).fillna(0).groupby(todos["final"]).sum()
    productos["_ns_anterior"] = productos["_uid"].map(anterior).fillna(0)
    ids = ids[ids["Identificador interno"].isin(ns.index)].drop_duplicates("Identificador interno")
    # Artículos de NetSuite con el mismo código de barras que otro (hoja Duplicados de
    # NetSuite) se suman al registro que tiene ese código.
    fila_cb = {cb.lstrip("0"): u for cb, u in zip(productos["Código de barras"], productos["_uid"]) if cb}
    sueltos = ns[~ns.index.isin(ids["Identificador interno"])]
    sueltos = sueltos[sueltos["Código de barras"].str.lstrip("0").isin(fila_cb)]
    ids = pd.concat([ids[["final", "Identificador interno"]], pd.DataFrame({
        "final": sueltos["Código de barras"].str.lstrip("0").map(fila_cb).to_numpy(),
        "Identificador interno": sueltos.index})], ignore_index=True)
    grupos = ids.groupby("final")["Identificador interno"].agg(list)
    fila = dict(zip(productos["_uid"], productos.index))
    for uid, lista in grupos.items():
        if uid not in fila:
            continue
        i = fila[uid]
        datos = ns.loc[lista]
        productos.at[i, "Stock Sistema NetSuite"] = f"{numero(datos['Stock Sistema NetSuite']).sum():g}"
        productos.at[i, "Disponible NetSuite"] = f"{numero(datos['Disponible NetSuite']).sum():g}"
        for c in ("Ubicaciones con stock (NetSuite)", "Ubicaciones con existencia negativa"):
            productos.at[i, c] = "; ".join(x for x in datos[c] if x)
        productos.at[i, "¿Alguna ubicación en negativo?"] = \
            "Sí" if (datos["¿Alguna ubicación en negativo?"] == "Sí").any() else "No"


def marcar_shopify_ariat(productos, shopify):
    """Agrega la tienda Shopify Ariat a las columnas de plataformas de Productos."""
    con = shopify[(shopify["En mano"] != 0) & shopify["uid"].notna()]
    texto = (con["Ubicación"] + ": " + con["En mano"].map("{:g}".format)).groupby(con["uid"]).agg(" | ".join)
    for uid, valor in texto.items():
        i = productos.index[productos["_uid"] == uid]
        if not len(i):
            continue
        i = i[0]
        actual = productos.at[i, "📍 Ubicación en plataformas (Plataforma + Marca)"]
        if "Shopify Ariat México" not in actual:
            productos.at[i, "📍 Ubicación en plataformas (Plataforma + Marca)"] = \
                " | ".join(x for x in (actual, "Shopify Ariat México") if x)
        actual = productos.at[i, "Ubicaciones con stock (plataformas)"]
        productos.at[i, "Ubicaciones con stock (plataformas)"] = " | ".join(x for x in (actual, valor) if x)


# --------------------------------------------------------------------------
# Stock por Ubicación (mismas columnas que la plantilla)
# --------------------------------------------------------------------------

def stock_por_ubicacion(detalle, ciclico_stock, shopify):
    """NetSuite al corte actual, Shopify Ariat y las demás plataformas del Cíclico."""
    ns = detalle[detalle["En mano"] != 0]
    ns = pd.DataFrame({
        "Ubicación (NetSuite / Plataforma)": ns["Ubicación corta"],
        "Código de barras": ns["Código de barras"], "WB SKU": ns["WB SKU"],
        "Identificador interno": ns["Identificador interno"], "WB Marca": ns["Marca"],
        "WB Descripción": ns["Descripción"], "Stock en ubicación (Físico)": ns["En mano"],
        "Origen en Productos": f"NetSuite {CORTE_NS}", "Fila en Productos": "",
        "Sistema": "NetSuite"})
    otras = ciclico_stock[ciclico_stock["Sistema"] != "NetSuite"].copy()
    sa = shopify[shopify["En mano"] != 0]
    sa = pd.DataFrame({
        "Ubicación (NetSuite / Plataforma)": sa["Ubicación"], "Código de barras": sa["Código de barras"],
        "WB SKU": sa["WB SKU"], "Identificador interno": sa["Variant ID"], "WB Marca": "Ariat",
        "WB Descripción": sa["Descripción"], "Stock en ubicación (Físico)": sa["En mano"],
        "Origen en Productos": f"Shopify Ariat {CORTE_SHOPIFY_ARIAT}", "Fila en Productos": "",
        "Sistema": "Shopify", "uid": sa["uid"]})
    tabla = pd.concat([ns.assign(uid=detalle.loc[ns.index, "uid"]), otras, sa], ignore_index=True)
    tabla["Stock en ubicación (Físico)"] = numero(tabla["Stock en ubicación (Físico)"])
    return tabla


# --------------------------------------------------------------------------
# Hoja Inventario
# --------------------------------------------------------------------------

IDENTIFICACION = ["Código de barras", "Identificador interno", "WB SKU", "WB N.º de estilo",
                  "WB Marca", "WB Descripcion", "WB Talla", "WB Color", "WB División",
                  "WB Categoría", "WB Género", "Unidad", "Precio de compra", "Precio de venta"]


def tabla_inventario(variantes, detalle, anterior, plataformas, shopify, escaneos):
    """Un renglón por artículo de la marca con cualquier existencia o movimiento."""
    uids = set(variantes["_uid"])
    det = detalle[detalle["uid"].isin(uids)]
    tot = det.groupby("uid")[CANTIDADES].sum()
    inactivas = det[det["Ubicación activa"] == "No"].groupby("uid")["En mano"].sum()
    neg = det[det["En mano"] < 0]
    negativos = (neg["Ubicación corta"] + " (" + neg["En mano"].map("{:g}".format) + ")") \
        .groupby(neg["uid"]).agg("; ".join)
    por_ubicacion = det[det["En mano"] != 0].pivot_table(
        index="uid", columns="Ubicación corta", values="En mano", aggfunc="sum")
    plat = plataformas[plataformas["uid"].isin(uids)]
    plat = plat.assign(P=plat["Ubicación (NetSuite / Plataforma)"].str.split(" › ").str[0]) \
        .pivot_table(index="uid", columns="P", values="Stock en ubicación (Físico)", aggfunc="sum")
    sa = shopify[shopify["uid"].isin(uids)]
    sa_tot = sa.groupby("uid")[["En mano", "Disponible", "Por llegar"]].sum()
    sa_tienda = sa[sa["En mano"] != 0].pivot_table(index="uid", columns="Tienda",
                                                   values="En mano", aggfunc="sum")
    esc = escaneos[escaneos["uid"].isin(uids)].groupby("uid").agg(
        piezas=("Piezas Contadas", "sum"), fecha=("Fecha / Hora", "max"))

    t = variantes.set_index("_uid")[IDENTIFICACION].copy()
    for c in ("Precio de compra", "Precio de venta"):
        t[c] = pd.to_numeric(t[c], errors="coerce")
    for c in CANTIDADES:
        t[f"{c} NetSuite"] = tot[c].reindex(t.index).fillna(0)
    t["En mano en ubicaciones inactivas"] = inactivas.reindex(t.index).fillna(0)
    t["Ubicaciones en negativo"] = negativos.reindex(t.index).fillna("")
    t[f"En mano NetSuite {CORTE_ANTERIOR}"] = anterior.reindex(t.index).fillna(0)
    t["Variación neta"] = 0.0  # Fórmula en la hoja.
    if len(sa_tot):
        t["Shopify Ariat en mano"] = sa_tot["En mano"].reindex(t.index).fillna(0)
        t["Shopify Ariat disponible"] = sa_tot["Disponible"].reindex(t.index).fillna(0)
        t["Shopify Ariat por llegar"] = sa_tot["Por llegar"].reindex(t.index).fillna(0)
    for p in PLATAFORMAS:
        if p in plat.columns:
            t[p] = plat[p].reindex(t.index).fillna(0)
    t["Piezas contadas (Escaneos)"] = esc["piezas"].reindex(t.index).fillna(0) if len(esc) else 0
    t["Fecha del último conteo"] = esc["fecha"].reindex(t.index).fillna("") if len(esc) else ""
    t["Valor en mano a costo"] = 0.0
    t["Valor en mano a precio de venta"] = 0.0
    cantidades = [c for c in t.columns if t[c].dtype.kind in "fi"
                  and c not in ("Precio de compra", "Precio de venta", "Variación neta",
                                "Valor en mano a costo", "Valor en mano a precio de venta")]
    en_mano = t["En mano NetSuite"]
    t["Estado del inventario"] = [
        "Existencia negativa" if m < 0 else "Con existencia" if m > 0 else
        "Solo en plataformas" if any(t.at[u, c] for c in PLATAFORMAS + ["Shopify Ariat en mano"]
                                     if c in t.columns) else
        "En pedido o tránsito" if t.at[u, "En pedido NetSuite"] or t.at[u, "En tránsito NetSuite"] else
        "Contado en rack, sin existencia en NetSuite" if t.at[u, "Piezas contadas (Escaneos)"] else
        f"Sin existencia (tenía al {CORTE_ANTERIOR})" if t.at[u, f"En mano NetSuite {CORTE_ANTERIOR}"] else
        "Revisar: disponible o comprometido sin existencia" for u, m in en_mano.items()]
    ubicaciones = list(por_ubicacion.columns)
    for u in ubicaciones:
        t[f"NS · {u}"] = por_ubicacion[u].reindex(t.index).fillna(0)
    tiendas = list(sa_tienda.columns)
    for u in tiendas:
        t[f"Shopify Ariat · {u}"] = sa_tienda[u].reindex(t.index).fillna(0)
    # Inventario actual: fuera quedan los artículos cuya única señal es una orden de compra
    # abierta (en pedido); sus unidades se reportan aparte en los KPIs.
    sin_pedido = t[[c for c in cantidades if c not in ("En mano NetSuite", "En pedido NetSuite")]] \
        .abs().sum(axis=1)
    activo = (en_mano != 0) | (sin_pedido != 0)
    pedido = {"unidades": float(t["En pedido NetSuite"].sum()),
              "articulos": int((t["En pedido NetSuite"] != 0).sum()),
              "solo_pedido": int((~activo & (t["En pedido NetSuite"] != 0)).sum())}
    t = t[activo]
    t = t.sort_values(["WB División", "WB Categoría", "WB N.º de estilo", "WB Talla"])
    return t.reset_index(drop=True), ubicaciones, tiendas, pedido


def movimientos_por_ubicacion(detalle, anterior_ubicacion, uids):
    """Entradas y salidas netas por ubicación entre el corte anterior y el actual."""
    ahora = detalle[detalle["uid"].isin(uids)].groupby(["uid", "Ubicación corta"])["En mano"].sum()
    antes = anterior_ubicacion[anterior_ubicacion["uid"].isin(uids)].groupby(
        ["uid", "Ubicación"])["Stock"].sum()
    antes.index = antes.index.set_names(["uid", "Ubicación corta"])
    delta = ahora.sub(antes, fill_value=0).reset_index(name="d")
    tabla = delta.groupby("Ubicación corta")["d"].agg(
        entradas=lambda s: s[s > 0].sum(), salidas=lambda s: s[s < 0].sum())
    tabla["antes"] = antes.groupby(level=1).sum().reindex(tabla.index).fillna(0)
    return tabla


# --------------------------------------------------------------------------
# Escritura de las dos hojas nuevas
# --------------------------------------------------------------------------

def letra(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def escribir_inventario(wb, tabla, ubicaciones, tiendas, pedido, ubic_info, movimientos, grupo):
    from openpyxl.styles import Alignment, Font, PatternFill

    titulo = Font(bold=True, size=14)
    negrita = Font(bold=True)
    encabezado = PatternFill("solid", fgColor="1F3864")
    blanco = Font(bold=True, color="FFFFFF")
    seccion = PatternFill("solid", fgColor="D9E1F2")
    moneda, entero = '"$"#,##0.00', "#,##0"

    for nombre in ("Inventario", "KPIs Inventario"):  # La plantilla puede traerlas de una corrida anterior.
        if nombre in wb.sheetnames:
            del wb[nombre]
    ws = wb.create_sheet("Inventario")
    columnas = list(tabla.columns)
    col = {c: letra(i) for i, c in enumerate(columnas, start=1)}
    n = len(tabla) + 1
    for j, c in enumerate(columnas, start=1):
        celda = ws.cell(1, j, c)
        celda.font, celda.fill = blanco, encabezado
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 45
    numericas = {c for c in columnas if tabla[c].dtype.kind in "fi"}
    for i, fila in enumerate(tabla.itertuples(index=False), start=2):
        for j, (c, v) in enumerate(zip(columnas, fila), start=1):
            if c == "Variación neta":
                v = f"={col['En mano NetSuite']}{i}-{col[f'En mano NetSuite {CORTE_ANTERIOR}']}{i}"
            elif c == "Valor en mano a costo":
                v = f"=IFERROR({col['En mano NetSuite']}{i}*{col['Precio de compra']}{i},0)"
            elif c == "Valor en mano a precio de venta":
                v = f"=IFERROR({col['En mano NetSuite']}{i}*{col['Precio de venta']}{i},0)"
            elif c in numericas:
                v = None if pd.isna(v) else (float(v) if v % 1 else int(v))
                if v == 0 and c not in ("En mano NetSuite",):
                    v = None  # Celda vacía = 0, como en NetSuite; la hoja queda más ligera.
            elif v == "":
                v = None
            celda = ws.cell(i, j, v)
            if c.startswith(("Valor", "Precio")):
                celda.number_format = moneda
    for j, c in enumerate(columnas, start=1):
        ws.column_dimensions[letra(j)].width = 14 if c in numericas else 22
    ws.column_dimensions[col["WB Descripcion"]].width = 40
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{letra(len(columnas))}{max(n, 2)}"

    # KPIs: fórmulas sobre la hoja Inventario (se recalculan si se edita en Sheets).
    k = wb.create_sheet("KPIs Inventario")
    r = lambda c: f"Inventario!${col[c]}$2:${col[c]}${max(n, 2)}"
    k["A1"] = f"Indicadores de inventario · {grupo}"
    k["A1"].font = titulo
    k["A2"] = (f"NetSuite al {CORTE_NS} · corte anterior NetSuite {CORTE_ANTERIOR} · Shopify Ariat "
               f"{CORTE_SHOPIFY_ARIAT} · Shopify Stetson, Western Brothers y Odoo {CORTE_PLATAFORMAS}")
    fila = 4

    def bloque(nombre, encabezados):
        nonlocal fila
        k.cell(fila, 1, nombre).font = negrita
        for j in range(1, max(len(encabezados), 4) + 1):
            k.cell(fila, j).fill = seccion
        fila += 1
        for j, h in enumerate(encabezados, start=1):
            c = k.cell(fila, j, h)
            c.font, c.fill = blanco, encabezado
        fila += 1

    en_mano = "En mano NetSuite"
    bloque("Existencia actual (NetSuite)", ["Indicador", "Valor", "Nota"])
    indicadores = [
        ("Artículos con existencia", f'=COUNTIF({r(en_mano)},">0")', entero, ""),
        ("Unidades en mano", f"=SUM({r(en_mano)})", entero, "Existencia del sistema"),
        ("Unidades disponibles", f"=SUM({r('Disponible NetSuite')})", entero, ""),
        ("Unidades comprometidas", f"=SUM({r('Comprometido NetSuite')})", entero, "Pedidos por surtir"),
        ("Unidades en tránsito", f"=SUM({r('En tránsito NetSuite')})", entero, ""),
        ("Unidades en pedido", pedido["unidades"], entero,
         f"Órdenes de compra abiertas de {pedido['articulos']} artículos; {pedido['solo_pedido']} de "
         "ellos aún no tienen existencia y no aparecen en la hoja Inventario"),
        ("Unidades pendientes por surtir", f"=SUM({r('Pendiente por surtir NetSuite')})", entero, ""),
        ("Unidades en ubicaciones inactivas", f"=SUM({r('En mano en ubicaciones inactivas')})", entero,
         "Ubicaciones dadas de baja en NetSuite que aún tienen existencia"),
        ("Artículos con existencia negativa", f'=COUNTIF({r(en_mano)},"<0")', entero, "Revisar en conteo"),
        ("Unidades en negativo", f'=SUMIF({r(en_mano)},"<0")', entero, ""),
        ("Valor en mano a costo", f"=SUM({r('Valor en mano a costo')})", moneda,
         "Precio de compra de NetSuite; hay costos que parecen estar en dólares (ver Notas)"),
        ("Valor en mano a precio de venta", f"=SUM({r('Valor en mano a precio de venta')})", moneda,
         "Precio de venta sin IVA"),
    ]
    for nombre, formula, formato, nota in indicadores:
        k.cell(fila, 1, nombre)
        c = k.cell(fila, 2, formula)
        c.number_format = formato
        k.cell(fila, 3, nota)
        fila += 1
    fila += 1

    bloque(f"Movimiento neto {CORTE_ANTERIOR} → {CORTE_NS}", ["Indicador", "Valor", "Nota"])
    var = "Variación neta"
    for nombre, formula, nota in [
            (f"Unidades en mano al {CORTE_ANTERIOR}", f"=SUM({r(f'En mano NetSuite {CORTE_ANTERIOR}')})", ""),
            ("Variación neta de unidades", f"=SUM({r(var)})", "En mano actual − en mano anterior"),
            ("Unidades que salieron (neto)", f'=SUMIF({r(var)},"<0")', "Ventas, traspasos o ajustes"),
            ("Unidades que entraron (neto)", f'=SUMIF({r(var)},">0")', "Recepciones, traspasos o ajustes"),
            ("Artículos que bajaron", f'=COUNTIF({r(var)},"<0")', ""),
            ("Artículos que subieron", f'=COUNTIF({r(var)},">0")', ""),
            ("Artículos que se agotaron", f'=COUNTIFS({r(en_mano)},0,{r(f"En mano NetSuite {CORTE_ANTERIOR}")},">0")', "")]:
        k.cell(fila, 1, nombre)
        k.cell(fila, 2, formula).number_format = entero
        k.cell(fila, 3, nota)
        fila += 1
    fila += 1

    plataformas = [c for c in ["Shopify Ariat en mano"] + PLATAFORMAS if c in col]
    if plataformas:
        bloque("Existencia en plataformas", ["Plataforma", "Unidades", "Artículos con existencia", "Corte"])
        for p in plataformas:
            k.cell(fila, 1, p.replace(" en mano", ""))
            k.cell(fila, 2, f"=SUM({r(p)})").number_format = entero
            k.cell(fila, 3, f'=COUNTIF({r(p)},">0")').number_format = entero
            k.cell(fila, 4, CORTE_SHOPIFY_ARIAT if "Ariat" in p else CORTE_PLATAFORMAS)
            fila += 1
        fila += 1

    if "Piezas contadas (Escaneos)" in col:
        bloque("Conteo cíclico", ["Indicador", "Valor"])
        k.cell(fila, 1, "Piezas contadas en rack (Escaneos)")
        k.cell(fila, 2, f"=SUM({r('Piezas contadas (Escaneos)')})").number_format = entero
        fila += 1
        k.cell(fila, 1, "Artículos contados")
        k.cell(fila, 2, f'=COUNTIF({r("Piezas contadas (Escaneos)")},">0")').number_format = entero
        fila += 2

    bloque("Por ubicación NetSuite", ["Ubicación", "Bodega / tipo", "Activa", "En mano",
                                      "% del total", "Artículos con existencia", "Artículos en negativo",
                                      f"En mano {CORTE_ANTERIOR}", "Entradas netas", "Salidas netas",
                                      "Variación"])
    inicio = fila
    for u in ubicaciones:
        info = ubic_info.get(u, {})
        c = f"NS · {u}"
        mov = movimientos.loc[u] if u in movimientos.index else None
        k.cell(fila, 1, u)
        k.cell(fila, 2, info.get("tipo", ""))
        k.cell(fila, 3, info.get("activa", ""))
        k.cell(fila, 4, f"=SUM({r(c)})").number_format = entero
        k.cell(fila, 5, f"=IFERROR(D{fila}/SUM($D${inicio}:$D${inicio + len(ubicaciones) - 1}),0)").number_format = "0.0%"
        k.cell(fila, 6, f'=COUNTIF({r(c)},">0")').number_format = entero
        k.cell(fila, 7, f'=COUNTIF({r(c)},"<0")').number_format = entero
        k.cell(fila, 8, float(mov["antes"]) if mov is not None else 0).number_format = entero
        k.cell(fila, 9, float(mov["entradas"]) if mov is not None else 0).number_format = entero
        k.cell(fila, 10, float(mov["salidas"]) if mov is not None else 0).number_format = entero
        k.cell(fila, 11, f"=D{fila}-H{fila}").number_format = entero
        fila += 1
    otras = [u for u in movimientos.index if u not in ubicaciones]
    for u in otras:  # Ubicaciones que hoy quedaron en cero pero tenían existencia.
        mov = movimientos.loc[u]
        k.cell(fila, 1, u)
        k.cell(fila, 3, ubic_info.get(u, {}).get("activa", ""))
        k.cell(fila, 4, 0)
        k.cell(fila, 8, float(mov["antes"])).number_format = entero
        k.cell(fila, 9, float(mov["entradas"])).number_format = entero
        k.cell(fila, 10, float(mov["salidas"])).number_format = entero
        k.cell(fila, 11, f"=D{fila}-H{fila}").number_format = entero
        fila += 1
    fila += 1

    if tiendas:
        bloque(f"Tiendas Ariat: NetSuite {CORTE_NS} contra Shopify Ariat {CORTE_SHOPIFY_ARIAT}",
               ["Tienda NetSuite", "Tienda Shopify", "En mano NetSuite", "En mano Shopify", "Diferencia"])
        for ns_t, sh_t in TIENDAS_ARIAT.items():
            if f"NS · {ns_t}" not in col or f"Shopify Ariat · {sh_t}" not in col:
                continue
            k.cell(fila, 1, ns_t)
            k.cell(fila, 2, sh_t)
            k.cell(fila, 3, f"=SUM({r(f'NS · {ns_t}')})").number_format = entero
            k.cell(fila, 4, f"=SUM({r(f'Shopify Ariat · {sh_t}')})").number_format = entero
            k.cell(fila, 5, f"=D{fila}-C{fila}").number_format = entero
            fila += 1
        k.cell(fila, 1, "Pareja de tiendas por nombre; confirmar que correspondan.")
        fila += 2

    for campo in ("WB División", "WB Categoría"):
        valores = sorted(v for v in tabla[campo].unique() if v) + ([""] if (tabla[campo] == "").any() else [])
        bloque(f"Por {campo.replace('WB ', '').lower()}", [campo.replace("WB ", ""), "En mano",
                                                         "Artículos con existencia", "Valor a costo",
                                                         "Valor a precio de venta", "Variación neta"])
        for v in valores:
            criterio = f'"{v}"' if v else '""'
            k.cell(fila, 1, v or "(sin capturar)")
            k.cell(fila, 2, f"=SUMIFS({r(en_mano)},{r(campo)},{criterio})").number_format = entero
            k.cell(fila, 3, f'=COUNTIFS({r(campo)},{criterio},{r(en_mano)},">0")').number_format = entero
            k.cell(fila, 4, f"=SUMIFS({r('Valor en mano a costo')},{r(campo)},{criterio})").number_format = moneda
            k.cell(fila, 5, f"=SUMIFS({r('Valor en mano a precio de venta')},{r(campo)},{criterio})").number_format = moneda
            k.cell(fila, 6, f"=SUMIFS({r(var)},{r(campo)},{criterio})").number_format = entero
            fila += 1
        fila += 1

    # Listas fijas al corte: los 15 artículos con más existencia y con más salida.
    tabla_var = tabla.assign(_v=tabla[en_mano] - tabla[f"En mano NetSuite {CORTE_ANTERIOR}"])
    for nombre, datos in (("15 artículos con más unidades en mano", tabla_var.nlargest(15, en_mano)),
                          (f"15 artículos con más salida neta desde {CORTE_ANTERIOR}",
                           tabla_var[tabla_var["_v"] < 0].nsmallest(15, "_v"))):
        bloque(nombre, ["Código de barras", "WB SKU", "WB Descripcion", "WB Talla", "En mano", "Variación"])
        for _, x in datos.iterrows():
            for j, v in enumerate([x["Código de barras"], x["WB SKU"], x["WB Descripcion"], x["WB Talla"],
                                   float(x[en_mano]), float(x["_v"])], start=1):
                k.cell(fila, j, v)
            fila += 1
        fila += 1

    k.column_dimensions["A"].width = 42
    for c in "BCDEFGHIJK":
        k.column_dimensions[c].width = 18
    k.column_dimensions["C"].width = 30
