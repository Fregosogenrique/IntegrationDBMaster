"""Hoja «Anatomía de Productos» de cada Cíclico por marca.

Ordena todos los artículos de la marca con la estructura de
Anatomia_productos_atributos_G1522.md:

  · Padre o modelo: se identifica por WB Marca y WB N.º de estilo y lleva los atributos
    compartidos (marca, descripción, división, categoría, género, temporada, ciclo de vida
    y características compartidas).
  · Variante vendible: se identifica por WB SKU y Código de barras y lleva talla, talla de
    EE. UU., color, unidad, precios y los datos propios de la variante.
  · Unidad física: no se individualiza; se reportan las piezas de cada variante.

Un atributo de padre que cambia entre variantes no se sobrescribe por herencia: el padre
lo deja vacío y cada variante conserva su valor, con el caso anotado en la revisión.
homologar_ciclico_g1522.py llama a estas funciones al escribir cada Cíclico.
"""

import pandas as pd

HOJA = "Anatomía de Productos"

# Diccionario de atributos de la Anatomía: nivel de cada columna de Productos.
PADRE = ["Marca principal", "WB N.º de estilo", "WB Marca", "WB Descripcion", "WB División",
         "WB Categoría", "WB Género", "WB Licencia", "WB Temporada", "WB Ciclo de vida",
         "WB Fit", "WB Silueta", "WB Composición", "WB Descripción larga", "WB País de origen",
         "SAT Clave de producto", "Enlace de imagen"]
VARIANTE = ["Código de barras", "WB SKU", "Identificador interno", "WB Talla",
            "WB Talla de EE. UU.", "WB Color", "Unidad", "Precio de compra", "Precio de venta"]
# Atributos que identifican al modelo: si cambian entre sus variantes, el modelo está mal
# agrupado o capturado. En los demás la diferencia es válida y se conserva en la variante.
IDENTIDAD_PADRE = ["Marca principal", "WB Marca", "WB División", "WB Categoría", "WB Género"]
REQUERIDOS_PADRE = ["WB Marca", "WB Descripcion", "WB División", "WB Categoría", "WB Género"]
CATEGORIAS_FIT = {"Ropa", "Denim"}  # Por división.
CATEGORIAS_SILUETA = {"Sombreros", "Botas", "Zapatos", "Jeans"}

# Orígenes del Cíclico cuyo «Identificador interno» es un ID de NetSuite.
ORIGEN_NETSUITE = ("Catálogo existente", "Nuevo – NetSuite", "Hoja a Revisar")
ALTA = "¿Dado de alta en NetSuite?"
ID_NETSUITE = "ID interno NetSuite"

COLUMNAS = (["Nivel", "Modelo (padre)", "Variantes del modelo", ALTA, ID_NETSUITE] + PADRE[:7]
            + ["Familia de talla"] + PADRE[7:] + VARIANTE
            + ["Piezas en NetSuite", "Piezas en plataformas", "Revisión de anatomía"])
NUMERICAS = {"Variantes del modelo", "Precio de compra", "Precio de venta", "Piezas en NetSuite",
             "Piezas en plataformas"}


def numero(texto):
    try:
        return float(texto)
    except (TypeError, ValueError):
        return 0.0


def piezas_plataformas(texto):
    """Suma de «Ubicaciones con stock (plataformas)»: «Plataforma › Ubicación: n | …»."""
    return sum(numero(p.rpartition(":")[2]) for p in texto.split(" | ") if ":" in p)


def indice_netsuite(catalogo):
    """Llaves de los artículos dados de alta en NetSuite: el catálogo estándar trae todos
    los artículos de inventario activos (con y sin existencia)."""
    activos = catalogo[catalogo["Identificador interno"] != ""]
    por_cb, por_sku = {}, {}
    for i, cb, sku in zip(activos["Identificador interno"], activos["Código de barras"], activos["WB SKU"]):
        if cb.lstrip("0"):
            por_cb.setdefault(cb.lstrip("0"), i)
        if sku:
            por_sku.setdefault(sku.upper(), i)
    return {"ids": set(activos["Identificador interno"]), "cb": por_cb, "sku": por_sku}


def alta_netsuite(r, netsuite):
    """(estado, ID de NetSuite) de una variante: por su ID interno; si no, por su código de
    barras (o uno alterno) o su SKU, que en NetSuite pueden estar en un artículo con otro ID."""
    ident = r["Identificador interno"]
    if ident in netsuite["ids"]:
        return "Sí", ident
    codigos = [r["Código de barras"]] + str(r.get("Códigos de barras alternos", "")).split(" | ")
    for cb in codigos:
        if cb.lstrip("0") in netsuite["cb"]:
            return "Sí, por código de barras", netsuite["cb"][cb.lstrip("0")]
    if r["WB SKU"].upper() in netsuite["sku"]:
        return "Sí, por SKU", netsuite["sku"][r["WB SKU"].upper()]
    if ident and r["Origen del registro"].startswith(ORIGEN_NETSUITE):
        return "No: ya no está activo en NetSuite", ""
    return "No", ""


def tabla_anatomia(variantes, familias_de, criterios, netsuite=None):
    """Renglones de la hoja: cada padre seguido de sus variantes.

    «familias_de(categoría, género, criterios)» da las familias de talla válidas
    (Criterios › 3 y 8); «netsuite» (indice_netsuite) dice qué variantes están dadas de
    alta en NetSuite."""
    v = variantes.reset_index(drop=True).copy()
    sin_estilo = v["WB N.º de estilo"] == ""
    # Sin estilo no hay llave de padre: se agrupa por la descripción del modelo y se anota.
    v["_modelo"] = v["WB Marca"] + " · " + v["WB N.º de estilo"].where(
        ~sin_estilo, "Sin estilo · " + v["WB Descripcion"].where(v["WB Descripcion"] != "", v["WB SKU"]))
    v["_orden"] = range(len(v))
    if netsuite is not None:
        altas = [alta_netsuite(r, netsuite) for _, r in v.iterrows()]
        v["_alta"] = [a for a, _ in altas]
        v["_id_ns"] = [i for _, i in altas]
    else:
        v["_alta"] = v["_id_ns"] = ""
    division_de = {c: d for c, d in criterios.get("divisiones", {}).items()}

    padres = []
    for modelo, grupo in v.groupby("_modelo", sort=False):
        comun, distintos = {}, []
        for c in PADRE:
            valores = grupo[c].unique()
            if len(valores) == 1:
                comun[c] = valores[0]
            else:
                comun[c] = ""
                distintos.append(c)
        padres.append((modelo, grupo, comun, distintos))
    padres.sort(key=lambda p: (p[2]["WB División"] or "~", p[2]["WB Categoría"] or "~", p[0].lower()))

    filas = []
    for modelo, grupo, comun, distintos in padres:
        revision = []
        if grupo["WB N.º de estilo"].iloc[0] == "":
            revision.append("Sin WB N.º de estilo: variantes agrupadas por descripción")
        faltan = [c for c in REQUERIDOS_PADRE if c not in distintos and not comun[c]]
        if faltan:
            revision.append("Falta: " + ", ".join(faltan))
        identidad = [c for c in distintos if c in IDENTIDAD_PADRE]
        if identidad:
            revision.append("Atributo de modelo distinto entre variantes: " + ", ".join(identidad))
        propios = [c for c in distintos if c not in IDENTIDAD_PADRE]
        if propios:
            revision.append("Cambia entre variantes, se conserva en cada una: " + ", ".join(propios))
        categoria, division = comun["WB Categoría"], comun["WB División"]
        if categoria and division and division_de.get(categoria, division) != division:
            revision.append(f"División no corresponde a {categoria} ({division_de[categoria]})")
        if comun["WB Fit"] and division and division not in CATEGORIAS_FIT:
            revision.append(f"Fit no aplica en {division}")
        if comun["WB Silueta"] and categoria and categoria not in CATEGORIAS_SILUETA:
            revision.append(f"Silueta no aplica en {categoria}")
        familias = sorted(familias_de(categoria, comun["WB Género"], criterios)) if categoria else []
        familia = " · ".join(familias) or criterios["familias"].get(categoria, "")
        nacional = sum(numero(x) for x in grupo["Stock Sistema NetSuite"])
        plataformas = sum(piezas_plataformas(x) for x in grupo["Ubicaciones con stock (plataformas)"])

        dadas = int(grupo["_alta"].str.startswith("Sí").sum())
        alta = "" if netsuite is None else ("Sí" if dadas == len(grupo) else "No" if not dadas
                                            else f"Parcial: {dadas} de {len(grupo)} variantes")
        padre = {"Nivel": "Padre", "Modelo (padre)": modelo, "Variantes del modelo": len(grupo),
                 ALTA: alta,
                 "Familia de talla": familia, "Piezas en NetSuite": nacional,
                 "Piezas en plataformas": plataformas, "Revisión de anatomía": " · ".join(revision)}
        padre.update(comun)
        filas.append(padre)
        for _, r in grupo.sort_values("_orden").iterrows():
            variante = {"Nivel": "Variante", "Modelo (padre)": modelo, ALTA: r["_alta"],
                        ID_NETSUITE: r["_id_ns"]}
            variante.update({c: r[c] for c in distintos})  # La diferencia se queda en la variante.
            variante.update({c: r[c] for c in VARIANTE})
            for c in ("Precio de compra", "Precio de venta"):
                variante[c] = numero(r[c]) if r[c] else ""
            variante["Piezas en NetSuite"] = numero(r["Stock Sistema NetSuite"])
            variante["Piezas en plataformas"] = piezas_plataformas(r["Ubicaciones con stock (plataformas)"])
            avisos = []
            if not r["WB Talla"]:
                avisos.append("Sin WB Talla (Talla única si no tiene talla física)")
            incidencias = r.get("Incidencias G1522", "")
            if incidencias:
                avisos.append(incidencias)
            variante["Revisión de anatomía"] = " · ".join(avisos)
            filas.append(variante)
    return pd.DataFrame(filas, columns=COLUMNAS).fillna("")


def escribir_anatomia(wb, tabla):
    """Agrega la hoja al libro con el formato de Productos de la plantilla (Arial 10,
    encabezado azul marino, código de barras como texto, precios en pesos): cada padre y,
    debajo, sus variantes agrupadas. Las piezas del padre son la suma de sus variantes."""
    from copy import copy

    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter as letra

    if HOJA in wb.sheetnames:
        del wb[HOJA]
    productos = wb["Productos"]
    columna_de = {productos.cell(1, j).value: j for j in range(1, productos.max_column + 1)}
    texto = productos.cell(2, columna_de["WB Descripcion"])._style
    cifra = productos.cell(2, columna_de["Stock Sistema NetSuite"])._style
    estilo = []
    for c in COLUMNAS:
        if c == ID_NETSUITE:
            estilo.append(productos.cell(2, columna_de["Identificador interno"])._style)
        elif c in columna_de:
            estilo.append(productos.cell(2, columna_de[c])._style)
        else:
            estilo.append(cifra if c in NUMERICAS else texto)
    ws = wb.create_sheet(HOJA)
    ws.sheet_properties.tabColor = "1D3556"
    for j, c in enumerate(COLUMNAS, start=1):
        celda = ws.cell(1, j, c)
        celda._style = copy(productos.cell(1, 1)._style)
    ws.row_dimensions[1].height = productos.row_dimensions[1].height
    fondo_padre = PatternFill("solid", fgColor="DCE3EE")
    pieza = {j for j, c in enumerate(COLUMNAS) if c.startswith("Piezas")}
    fin = len(tabla) + 1
    niveles = tabla["Nivel"].tolist()
    for k, fila in enumerate(tabla.itertuples(index=False), start=2):
        padre = fila[0] == "Padre"
        if padre:  # Sus variantes ocupan los renglones siguientes hasta el próximo padre.
            ultimo = k
            while ultimo - 1 < len(niveles) and niveles[ultimo - 1] == "Variante":
                ultimo += 1
        for j, v in enumerate(fila):
            if padre and j in pieza and ultimo > k:
                v = f"=SUM({letra(j + 1)}{k + 1}:{letra(j + 1)}{ultimo})"
            celda = ws.cell(k, j + 1, None if v == "" else v)
            celda._style = copy(estilo[j])  # Copia: el estilo de origen no debe cambiar.
            if j in pieza or COLUMNAS[j] == "Variantes del modelo":
                celda.number_format = "#,##0"
            if padre:
                fuente = copy(celda.font)
                fuente.b = True
                celda.font, celda.fill = fuente, fondo_padre
        if not padre:
            ws.row_dimensions[k].outlineLevel = 1  # Variantes agrupadas bajo su padre.
    anchos = {"Nivel": 10, "Modelo (padre)": 34, "Variantes del modelo": 11, "WB Descripcion": 40,
              "WB Descripción larga": 50, "Enlace de imagen": 30, "Familia de talla": 30,
              "Revisión de anatomía": 60, "Código de barras": 16, "WB SKU": 24, ALTA: 24}
    for j, c in enumerate(COLUMNAS, start=1):
        ws.column_dimensions[letra(j)].width = anchos.get(c, 13 if c in NUMERICAS else 18)
    ws.sheet_properties.outlinePr.summaryBelow = False  # El padre va arriba de su grupo.
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{letra(len(COLUMNAS))}{max(fin, 2)}"


def resumen(tabla):
    padres = tabla[tabla["Nivel"] == "Padre"]
    variantes = tabla[tabla["Nivel"] == "Variante"]
    return {"padres": len(padres), "variantes": len(variantes),
            "padres_revision": int((padres["Revisión de anatomía"] != "").sum()),
            "en_netsuite": int(variantes[ALTA].str.startswith("Sí").sum())}


if __name__ == "__main__":  # Vista previa sobre un Cíclico ya generado.
    import sys
    from homologar_ciclico_g1522 import familias_de
    datos = pd.read_excel(sys.argv[1], sheet_name="Productos", dtype=str, engine="calamine").fillna("")
    criterios = {"familias": {}, "tallas": [], "divisiones": {}}
    t = tabla_anatomia(datos, familias_de, criterios)
    print(resumen(t))
    print(t.head(12).to_string(max_colwidth=30))
